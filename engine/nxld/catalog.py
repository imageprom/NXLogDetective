"""NXLD: каталог структуры отчёта — data/locale/<язык>/report_structure.json (задача #8).

Каталог задаёт файлы отчёта (номер, имя, порядок) и листы каждого файла: код, статус, вкладку, заголовок, подзаголовок и строку
оглавления. Оформление находит, что рисовать, по коду листа (sheets.CODES: код → имя, под которым лист строит код); вкладка,
заголовок и подзаголовок берутся из каталога.

Статус: approved — выводится; review — только с ключом запуска NXLD_SHOW_REVIEW=1 (модельные прогоны для согласования);
off — не выводится нигде. Лист не approved не пишется, в журнал запуска — «лист «…» не утверждён, не выводится».

Тексты (heading, subheading) — строка или массив строк, строка — массив частей: {"text": "…"} — статика из каталога,
{"auto": "код"} — динамика, которую подставляет оформление (содержимое целиком за ним). Части одной строки — через пробел,
строки подзаголовка — с новой строки; пустая часть и строка без частей не выводятся."""
import json
import os
import re

LOCALE = os.environ.get('NXLD_LOCALE', 'ru')
_PATH = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..', 'data', 'locale', LOCALE, 'report_structure.json'))
_CAT = None
CURRENT = None   # блок файла, который сейчас собирается


def load(path=None):
    """Каталог (кэшируется; path — для тестов и правок каталога без правки кода)."""
    global _CAT
    if path is not None:
        return json.load(open(path, encoding='utf-8'))
    if _CAT is None:
        p = os.environ.get('NXLD_CATALOG') or _PATH
        _CAT = json.load(open(p, encoding='utf-8'))
    return _CAT


def reset():
    global _CAT
    _CAT = None


def show_review():
    return os.environ.get('NXLD_SHOW_REVIEW', '') not in ('', '0')


def files():
    """Файлы отчёта в порядке каталога: [{num, code, file, title, audience, sheets}]."""
    return load()['files']


def file_entry(block):
    from .sheets import BLOCK_CODES
    code = BLOCK_CODES[block]
    return next(f for f in files() if f['code'] == code)


def block_of(code):
    from .sheets import BLOCK_CODES
    return next(b for b, c in BLOCK_CODES.items() if c == code)


def blocks():
    """Блоки в порядке файлов каталога."""
    return [block_of(f['code']) for f in files()]


def file_names():
    """{блок: имя файла} в порядке каталога (NXLD_01_Overview.xlsx …)."""
    return {block_of(f['code']): f['file'] for f in files()}


def sheets(block):
    return file_entry(block)['sheets']


def shown_status(st):
    return st == 'approved' or (st == 'review' and show_review())


def internal(block, code):
    """Имя, под которым лист строит код (sheets.CODES), по коду листа."""
    from .sheets import CODES
    return CODES[block][code]


def by_internal(block, name):
    """Запись каталога по имени, под которым лист строит код; None — листа в каталоге нет."""
    from .sheets import CODES
    code = next((c for c, n in CODES.get(block, {}).items() if n == name), None)
    return next((s for s in sheets(block) if s['code'] == code), None) if code else None


def shown(block, name):
    """Выводится ли лист (по имени в коде): в каталоге и утверждён (или review при NXLD_SHOW_REVIEW)."""
    e = by_internal(block, name)
    return bool(e) and shown_status(e['status'])


def order(block):
    """Имена (в коде) выводимых листов файла в порядке каталога."""
    return [internal(block, s['code']) for s in sheets(block) if shown_status(s['status'])]


def tab(block, name):
    """Вкладка листа по имени в коде (как в каталоге); вне каталога — имя как есть."""
    e = by_internal(block, name)
    return e['tab'] if e else name


def card_tab(block, name):
    """Вкладка для ссылки карточки: выводимый лист — его вкладка, невыводимый — Обзор своего файла."""
    return tab(block, name) if shown(block, name) else tab(block, 'Обзор')


def toc(block, name):
    e = by_internal(block, name)
    return e.get('toc', '') if e else ''


def render(text, auto=None):
    """Текст каталога → строка ячейки: строки — через перевод строки, части строки — через пробел; пустое не выводится."""
    if text is None: return None
    if isinstance(text, str): return text
    auto = auto or {}
    lines = []
    for line in text:
        parts = []
        for p in line:
            v = p.get('text') if 'text' in p else auto.get(p.get('auto'))
            if v not in (None, ''): parts.append(str(v))
        if parts: lines.append(' '.join(parts))
    return '\n'.join(lines)


# ---- регистрация заголовков при записи листа ----
def mark(ws, heading=None, sub=None, **auto):
    """Оформление сообщает, в каких ячейках заголовок и подзаголовок листа и какие у них динамические части (auto).
    Код пишет текст как раньше; finalize заменяет его текстом каталога."""
    ws._nxld_head = dict(heading=heading, sub=sub, auto=auto)


def prune(wb, block, log=print):
    """Убрать из книги листы, которые не выводятся (не approved), — до построения Обзора: строки Обзоров на них не появляются.
    Возвращает имена убранных листов."""
    from .sheets import SERVICE
    gone = []
    for ws in list(wb.worksheets):
        if ws.title in SERVICE: continue
        e = by_internal(block, ws.title)
        if e is not None and not shown_status(e['status']):
            log(f"лист «{e['tab']}» не утверждён, не выводится")
            del wb[ws.title]
            gone.append(ws.title)
    return gone


def missing(block, names):
    """Листы файла, которых нет в каталоге: сборка их создала, а каталог их не знает."""
    from .sheets import SERVICE, CODES
    known = set(CODES.get(block, {}).values()) | set(SERVICE)
    return [n for n in names if n not in known]


_LINK = re.compile(r"^(?P<file>[^#]*)#'(?P<sheet>(?:[^']|'')+)'!(?P<cell>\S+)$")


def target_map(block):
    """Имя листа в коде → вкладка в отчёте этого файла; невыводимые листы → вкладка Обзора файла."""
    from .sheets import CODES
    first = tab(block, 'Обзор')
    out = {}
    for code, name in CODES.get(block, {}).items():
        e = next((s for s in sheets(block) if s['code'] == code), None)
        out[name] = e['tab'] if e and shown_status(e['status']) else first
    return out


def finalize(wb, block, log=print):
    """Последний проход по книге файла: невыводимые листы убрать, заголовки и подзаголовки — из каталога, вкладки — из каталога,
    ссылки — на новые вкладки (на невыведенный лист — на Обзор своего файла), порядок — как в каталоге."""
    from .sheets import SERVICE
    prune(wb, block, log)
    for ws in wb.worksheets:   # заголовки и подзаголовки
        h = getattr(ws, '_nxld_head', None)
        e = by_internal(block, ws.title)
        if not h or not e: continue
        if h.get('heading') and e.get('heading') is not None:
            ws[h['heading']].value = render(e['heading'], h['auto'])
        if h.get('sub') and e.get('subheading') is not None:
            v = render(e['subheading'], h['auto'])
            ws[h['sub']].value = v if v else None
    # ссылки: внутри файла и на другие файлы этой проверки
    names = {block_of(f['code']): f['file'] for f in files()}
    maps = {f: target_map(b) for b, f in names.items()}
    here = maps[names[block]]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                hl = c.hyperlink
                if hl is None: continue
                loc = hl.location or hl.target or ''
                m_ = _LINK.match(loc) if '#' in loc else None
                if not m_: continue
                f_, sh_ = m_.group('file'), m_.group('sheet').replace("''", "'")
                mp = maps.get(f_) if f_ else here
                if mp is None or sh_ not in mp or mp[sh_] == sh_: continue
                new = mp[sh_]
                ref = (f"{f_}#'{new}'!A1" if f_ else f"#'{new}'!A1")
                c.hyperlink = ref
                if isinstance(c.value, str):
                    c.value = c.value.replace(f'«{sh_}»', f'«{new}»') if f'«{sh_}»' in c.value else (new if c.value == sh_ else c.value)
    # вкладки и порядок
    for ws in wb.worksheets:
        if ws.title in SERVICE: continue
        t = tab(block, ws.title)
        if t != ws.title: ws.title = t
    pos = {tab(block, n): i for i, n in enumerate(order(block))}
    svc = [w for w in wb._sheets if w.title in SERVICE]
    rest = [w for w in wb._sheets if w.title not in SERVICE]
    rest.sort(key=lambda w: pos.get(w.title, len(pos)))
    wb._sheets = rest + svc
    return wb


def fix_links(books, outdir=None):
    """Ссылки во всех файлах проверки — живые (#8). books — {имя файла: книга} собранных файлов; другие файлы папки — по диску.
    Карточка («Проблемы») со ссылкой на лист, которого нет, ведёт на Обзор файла этого листа; строка-ссылка на другом листе
    (Обзоры, «Связанное в других отчётах», ссылки под таблицами) на невыведенный лист не выводится: текст и ссылка убираются."""
    sheets_of = {f: set(wb.sheetnames) for f, wb in books.items()}
    if outdir:
        import glob
        from openpyxl import load_workbook
        for p in glob.glob(os.path.join(outdir, 'NXLD_*.xlsx')):
            f = os.path.basename(p)
            if f not in sheets_of:
                try: sheets_of[f] = set(load_workbook(p, read_only=True).sheetnames)
                except Exception: pass
    first = {f: next(iter(wb.sheetnames), None) for f, wb in books.items()}
    for f in sheets_of:
        first.setdefault(f, 'Обзор')
    fixed = 0
    for f, wb in books.items():
        for ws in wb.worksheets:
            card = ws.title == tab_by_code_any('problems')
            for row in ws.iter_rows():
                for c in row:
                    hl = c.hyperlink
                    if hl is None: continue
                    loc = hl.location or hl.target or ''
                    m_ = _LINK.match(loc) if '#' in loc else None
                    if not m_: continue
                    tf, sh = m_.group('file') or f, m_.group('sheet').replace("''", "'")
                    if tf in sheets_of and sh in sheets_of[tf]: continue
                    fixed += 1
                    if card and tf in sheets_of:   # карточка — на Обзор файла листа
                        ov = first[tf]
                        c.hyperlink = f"#'{ov}'!A1" if tf == f else f"{tf}#'{ov}'!A1"
                        if isinstance(c.value, str): c.value = c.value.replace(f'«{sh}»', f'«{ov}»')
                    else:   # строка-ссылка на невыведенный лист — не выводится
                        c.hyperlink = None
                        c.value = None
    return fixed


def tab_by_code_any(code):
    """Вкладка листа с этим кодом (одинакова во всех файлах, например «Проблемы»)."""
    for f in files():
        for s in f['sheets']:
            if s['code'] == code: return s['tab']
    return None
