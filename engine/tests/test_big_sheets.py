"""Большие листы: больше 10 000 строк — в Excel первые 1 000 и строка-пояснение, полный лист — CSV в архиве отчёта."""
import csv
import glob
import io
import os
import zipfile
from datetime import timedelta

import synth

N = 10500


def test_scanner_ips_cut_and_csv(tmp_path):
    log = synth.Log()
    log.people(per_day=10)
    for i in range(N):   # 10 500 разных IP сканеров: по два зонда с каждого
        ip = f'45.{140 + i // 65000}.{(i // 250) % 256}.{i % 250 + 1}'
        t = synth.T0 + timedelta(days=i % 7, hours=2, seconds=i)
        for k, u in enumerate(('/.env', '/wp-config.php.bak')):
            log.line(ip, t + timedelta(seconds=k), u, 404, ua='Mozilla/5.0 (Windows NT 6.1; rv:60.0) Gecko/20100101 Firefox/60.0')
    res, out, _ = synth.run(log, str(tmp_path))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out, 'NXLD_04_Bots.xlsx'))
    ws = wb['IP сканеров']
    note = ws['B3'].value
    assert note and note.startswith('Показаны 1 000 из ') and 'NXLD_04_Bots_scanner_ips.csv' in note
    hdr = [c.value for c in ws[4] if c.value is not None]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[1] and str(r[1]).count('.') == 3]   # строки таблицы — IP в колонке B
    assert len(rows) == 1000
    total = int(note.split(' из ')[1].split('.')[0].replace(' ', ''))
    assert total >= N
    z = zipfile.ZipFile(glob.glob(os.path.join(out, '*.zip'))[0])
    assert 'NXLD_04_Bots_scanner_ips.csv' in z.namelist()
    raw = z.read('NXLD_04_Bots_scanner_ips.csv')
    assert raw.startswith('﻿'.encode('utf-8'))
    R = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), delimiter=';'))
    assert R[0] == hdr
    assert len(R) - 1 == total
    assert [r[0] for r in R[1:1001]] == [r[1] for r in rows]   # Excel — первые строки CSV в той же сортировке


def test_raw_ip_sheet_cut(tmp_path):
    """Сырой лист «IP» (01, когда блок «Боты» не выбран) — та же обрезка в write_xlsx: пояснение над таблицей, CSV той же структуры."""
    import sys
    import pandas as pd
    sys.path.insert(0, synth.ENGINE)
    from nxld import report, report_tables
    df = pd.DataFrame({'ip': [f'10.0.{i // 250}.{i % 250}' for i in range(N)], 'запросов': range(N, 0, -1)})
    report_tables.BIG_OUT.update(файл='NXLD_01_Overview', csv=[])
    p = os.path.join(tmp_path, 'x.xlsx')
    report.write_xlsx(p, {'IP': df, 'Другое': df.head(5)})
    out = report_tables.write_big_csv(str(tmp_path))
    report_tables.BIG_OUT['файл'] = ''
    import openpyxl
    wb = openpyxl.load_workbook(p)
    assert wb['IP']['A1'].value.startswith('Показаны 1 000 из 10 500') and 'NXLD_01_Overview_IP.csv' in wb['IP']['A1'].value
    assert [c.value for c in wb['IP'][2]] == ['ip', 'запросов'] and wb['IP'].max_row == 1002
    assert wb['Другое'].max_row == 6
    assert [os.path.basename(f) for f in out] == ['NXLD_01_Overview_IP.csv']
    back = pd.read_csv(out[0], sep=';', encoding='utf-8-sig')
    assert list(back.columns) == ['ip', 'запросов'] and len(back) == N
