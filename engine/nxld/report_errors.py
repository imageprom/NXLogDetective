"""NXLD: файл 02 «Ошибки» — по тем же правилам, что Overview.

Обзор (шапка с логотипом и сводка блоками) → Проблемы → Статистика → вкладки по кодам (404, 5xx, 403 и прочие 4xx, 499)
→ Пострадавшие страницы → Файлы с ошибками → полные листы. Данные — из реестров (errors.py, classify, actors), без своих условий.
Подсветка везде одна: работает — белым, пострадавшее — персиковым, чужое и норма — серым."""
from . import sheets
from .catalog import order as catalog_order
import pandas as pd
from openpyxl.styles import Font
from .report_index import Wide, brand_header, GREY, ORANGE
from .report_tables import first_heading, data_sheet, journal_sheet, service_files, F_NOTE, _day, legend, fmt_codes, redden_codes, heat_sheet, extra_table
from .report_index import GREY as T_GREY
from .errors import WHO, FAMILIES, by_family, ERRLOG_MEANING

F_GREY = 'EFEFEF'
DROP = ('Коды ответа', 'Служебные файлы', 'Отсутствующие ресурсы', 'Ошибки по дням', '5xx по шаблонам', 'Битые ссылки',
        'Битые ссылки: страницы-источники', '404: входы извне', 'Изменения статусов', 'Идентичные ответы', 'Пустые страницы')   # последние два — лист «Мягкие ошибки»
RENAME = sheets.RENAME_02
# старые имена листов в карточках → новые
LINK = sheets.LINK_02


def relink(findings):
    for x in findings:
        if x.get('лист') in LINK: x['лист'] = LINK[x['лист']]


def sheets_for(S):
    """Что из расчётов блока пишется в файл: без замещённых листов, с новыми именами."""
    return {RENAME.get(k, k): v for k, v in S.items() if k not in DROP}


def _d(v):
    try: return pd.to_datetime(v).strftime('%d.%m.%Y')
    except Exception: return str(v or '')


def _t(v):
    try: return pd.to_datetime(v).strftime('%d.%m.%Y %H:%M')
    except Exception: return str(v or '')


def fill_of(alien, crit):
    """Общая подсветка 02: чужое — серой заливкой, критичное — персиковым, остальное — белым."""
    return F_GREY if alien else (F_NOTE if crit else None)


def font_of(actual, ill):
    """Неактуальное (починилось или давно не встречалось) — серым текстом."""
    return T_GREY if not actual and ill else None


from .report_tables import notes, notes_of, legend_of   # примечания и легенда — общие, тексты из каталога (#11)


def code_tab(wb, T, Sr, family):
    """Вкладка по коду — статистика: каждый адрес с этим кодом, кто получил, откуда пришли, что с ним сейчас и почему это важно.
    История адреса — на «Пострадавших страницах»; здесь её нет, чтобы не повторяться."""
    D_all = by_family(T, family, Sr)
    if not len(D_all): return None
    D = D_all.head(5000)
    D = D.sort_values('запросов', ascending=False)
    d = pd.DataFrame({'Адрес': D['адрес'], 'Текущий статус': D['сейчас'], 'Основание': D['почему'], 'Запросов': D['запросов'].astype(int),
                      **{nm: D[nm].astype(int) for nm in WHO}, 'Уникальных IP': D['IP'].astype(int), 'Ошибки': D['коды'].map(fmt_codes),
                      'Тип': D['тип'], 'Раздел': D['раздел'], 'Источники': D.get('источники', ''), 'Первый': D['первый'].map(_d), 'Последний': D['последний'].map(_d)})
    if family in ('404', '499'): d = d.drop(columns='Ошибки')   # код один — колонка ничего не добавляет
    fills = [fill_of(a, c) for a, c in zip(D['чужое'], D['критично'])]
    fonts = [None if a else font_of(ac, il) for a, ac, il in zip(D['чужое'], D['актуально'], D['болел'])]
    kpi = [('Адресов', len(D_all)), ('Критично', int(D_all['критично'].sum())), ('У людей', int(D_all['Люди'].sum())), ('Сканеры и чужое', int(D_all['чужое'].sum()))]
    widths = {'Адрес': 48, 'Запросов': 10, 'Уникальных IP': 10, 'Тип': 11, 'Текущий статус': 30, 'Основание': 26, 'Ошибки': 18, 'Раздел': 18, 'Источники': 46, 'Первый': 11, 'Последний': 11, **{nm: 10 for nm in WHO}}
    data_sheet(wb, family, d, widths, wrap=('Адрес', 'Текущий статус', 'Основание', 'Источники', 'Ошибки'), center=('Тип',),
               kpi=kpi, kpi_col='Мониторинг', row_rule=lambda r, _it=iter(fills): next(_it), font_rule=lambda r, _it=iter(fonts): next(_it), red=())
    notes_of(wb[family], cut='Показаны первые 5 000 адресов.' if len(D_all) > 5000 else '')
    legend_of(wb[family], ['критично', 'обычно', 'неактуально', 'чужое'])
    return wb[family]


def victims(wb, E):
    """«Пострадавшие страницы»: судьба каждой страницы сайта, на которой люди, свои или поисковики получали ошибку.
    «Сейчас» отдельно от «Истории»; критичное — персиковым, починившееся — серым текстом."""
    N = E.get('нерабочие')
    name = 'Пострадавшие страницы'
    if N is None or not len(N): return None
    if name not in wb.sheetnames: wb.create_sheet(name)
    N = N.sort_values('запросов', ascending=False).reset_index(drop=True)
    d = pd.DataFrame({'Страница': N['адрес'], 'Текущий статус': N['сейчас'], 'Основание': N['почему'], 'История': N['статус'], 'Запросов': N['запросов'].astype(int),
                      'Люди': N['Люди'].astype(int), 'Поисковики': N['Поисковики'].astype(int), 'Свои': N['Свои'].astype(int), 'Ошибки': N['коды'].map(fmt_codes),
                      'Раздел': N['раздел'], 'Источники': N.get('источники', ''), 'Первый': N['первый'].map(_d), 'Последний': N['последний'].map(_d)})
    fills = [fill_of(False, c) for c in N['критично']]
    fonts = [font_of(a, il) for a, il in zip(N['актуально'], N['болел'])]
    kpi = [('Страниц', len(d)), ('Критично', int(N['критично'].sum())), ('Не работают сейчас', int(N['актуально'].sum())), ('Починились', int((~N['актуально'] & N['болел']).sum()))]
    data_sheet(wb, name, d, {'Страница': 44, 'Текущий статус': 28, 'Основание': 28, 'История': 40, 'Ошибки': 20, 'Раздел': 18, 'Источники': 50},
               wrap=('Страница', 'Текущий статус', 'Основание', 'История', 'Источники', 'Ошибки'), kpi=kpi, kpi_col='Поисковики',
               row_rule=lambda r, _it=iter(fills): next(_it), font_rule=lambda r, _it=iter(fonts): next(_it), red=())
    notes_of(wb[name])
    legend_of(wb[name], ['критично', 'обычно', 'неактуально'])
    return wb[name]


def links_sheets(wb, E):
    """«Битые внутренние» и «Битые внешние» — по шаблонам, для отладки; персиковым — где ошибку получили люди."""
    I = E.get('битые_внутренние')
    if I is not None and len(I):
        nm = 'Битые внутренние'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Где стоит ссылка': I['откуда'], 'Куда ведёт': I['куда'], 'Переходов': I['переходов'].astype(int), 'Визитов': I['визитов'].astype(int),
                          'Люди': I['людей'].astype(int), 'Поисковики': I['поисковиков'].astype(int), 'Страниц-источников': I['страниц'].astype(int),
                          'Адресов-целей': I['адресов'].astype(int)})
        d.insert(2, 'Ошибки', I['коды'].map(fmt_codes))   # коды — перед переходами
        data_sheet(wb, nm, d, {'Где стоит ссылка': 42, 'Куда ведёт': 42, 'Ошибки': 20},
                   wrap=('Где стоит ссылка', 'Куда ведёт', 'Ошибки'), kpi=[('Пар шаблонов', len(d)), ('Переходов', int(d['Переходов'].sum()))], kpi_col='Люди',
                   row_rule=lambda r, _it=iter(list((I['людей'] > 0) & I['актуально'])): F_NOTE if next(_it) else None,
                   font_rule=lambda r, _it=iter(list(I['актуально'])): None if next(_it) else T_GREY, red=())
        notes_of(wb[nm])
        legend_of(wb[nm], ['люди', 'обычно', 'неактуально'])
    X = E.get('битые_внешние')
    if X is not None and len(X):
        nm = 'Битые внешние'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Откуда': X['вид'], 'Сайт': X['откуда'], 'Куда ведёт': X['куда'], 'Ошибки': X['коды'].map(fmt_codes), 'Переходов': X['переходов'].astype(int),
                          'Визитов': X['визитов'].astype(int), 'Люди': X['людей'].astype(int), 'Поисковики': X['поисковиков'].astype(int), 'Адресов-целей': X['адресов'].astype(int)})
        data_sheet(wb, nm, d, {'Сайт': 26, 'Куда ведёт': 46, 'Ошибки': 20}, center=('Откуда',),
                   wrap=('Куда ведёт', 'Ошибки'), kpi=[('Пар', len(d)), ('Переходов', int(d['Переходов'].sum()))], kpi_col='Люди',
                   row_rule=lambda r, _it=iter(list((X['людей'] > 0) & X['актуально'])): F_NOTE if next(_it) else None,
                   font_rule=lambda r, _it=iter(list(X['актуально'])): None if next(_it) else T_GREY, red=())
        notes_of(wb[nm])
        legend_of(wb[nm], ['люди', 'обычно', 'неактуально'])


def search_errors_sheet(wb, SE, nm='Поисковые ошибки'):
    """Поисковые ошибки — лист 06 «SEO» (или 02, если 06 не строится)."""
    if SE is None or not len(SE): return None
    if nm not in wb.sheetnames: wb.create_sheet(nm)
    d = pd.DataFrame({'Робот': SE['робот'], 'Тип страницы': SE['шаблон'], 'Ошибки': SE['коды'].map(fmt_codes), 'Запросов': SE['запросов'], 'Ошибок': SE['ошибок']})
    data_sheet(wb, nm, d, {'Тип страницы': 55, 'Ошибки': 24}, wrap=('Тип страницы', 'Ошибки'),
               row_rule=lambda r: F_NOTE if any(t.strip().startswith('5') for t in str(r.get('Ошибки', '')).split(',')) else None, red=())
    notes_of(wb[nm])
    legend_of(wb[nm], ['5xx', 'обычно'])
    redden_codes(wb, only=(nm,), rx=r'^Ошибки$')
    return wb[nm]


def soft_sheet(wb, G, P, nm='Мягкие ошибки'):
    """Мягкие ошибки: ответ 200, но не та страница — идентичные ответы на множестве адресов (основная таблица) и почти пустые ответы."""
    has_g, has_p = G is not None and len(G), P is not None and len(P)
    if not has_g and not has_p: return None
    if nm not in wb.sheetnames: wb.create_sheet(nm)
    def ident():
        return pd.DataFrame({'Размер, байт': G['размер'], 'Разных адресов': G['разных_адресов'].astype(int), 'Запросов': G['запросов'].astype(int),
                             'Людям, %': G['людям_%'], 'Поисковым роботам, %': G['поисковым_%'], 'Прочим роботам, %': G['прочим_%'], 'Ботам, %': G['ботам_%'],
                             'Примеры адресов': G['примеры'], 'Что это': G['что_это']})
    def empty():
        d = pd.DataFrame({'Адрес': P['base'].astype(str), 'Ответов меньше 300 байт': P['ответов_меньше_300_байт'].astype(int)})
        if 'что' in P: d['Что это'] = P['что']
        if 'комментарий' in P: d['Комментарий'] = P['комментарий']
        return d
    crit = lambda: [F_NOTE if (a >= 500 and (h > 0 or s_ > 0)) else None for a, h, s_ in zip(G['разных_адресов'], G['людям_%'], G['поисковым_%'])]
    norm = lambda: [F_GREY if x else None for x in (list(P['норма']) if 'норма' in P else [False] * len(P))]
    kpi = [('Групп идентичных ответов', len(G) if has_g else 0), ('Адресов в них', int(G['разных_адресов'].sum()) if has_g else 0), ('Пустых ответов', len(P) if has_p else 0)]
    from .catalog import K
    part = {'identical': (ident, crit, ['критично', 'обычно']), 'empty': (empty, norm, ['норма', 'обычно'])}
    have = [c for c in K(nm).order(['identical', 'empty']) if (has_g if c == 'identical' else has_p)]   # порядок таблиц — из каталога
    first = (have[0],) + part[have[0]]
    W_ = {'Размер, байт': 16, 'Примеры адресов': 46, 'Что это': 44, 'Адрес': 50, 'Комментарий': 55}
    WR = ('Примеры адресов', 'Что это', 'Адрес', 'Комментарий')
    data_sheet(wb, nm, first[1](), W_, wrap=WR, kpi=kpi, row_rule=lambda r, _it=iter(first[2]()): next(_it), red=())
    ws = wb[nm]
    first_heading(ws, first[0])   # заголовок первой таблицы — над шапкой
    for i, c in enumerate(have):
        df_, rule_, leg_ = part[c]
        if i: extra_table(ws, df_(), c, wrap=WR, row_rule=lambda r, _it=iter(rule_()): next(_it), red=())
        notes_of(ws, c)
        legend_of(ws, leg_, c)
    return ws


def full_sheets(wb, S, recent='', OUT=None):
    """Остальные полные листы — в едином табличном стиле, с человеческими колонками, подписями и подсветкой."""
    def put(name, df, rename, widths, wrap=(), kpi=None, kpi_col=None, rule=None, center=(), font=None):   # подзаголовок, ссылки, примечания — из каталога
        if df is None or not len(df) or name not in wb.sheetnames: return
        d = df.rename(columns=rename)[[v for v in rename.values() if v in df.rename(columns=rename).columns]]
        data_sheet(wb, name, d, widths, wrap=wrap, kpi=kpi, kpi_col=kpi_col, row_rule=rule, center=center, font_rule=font, red=())
    O = OUT
    if O is not None and len(O):
        O = O.assign(начало=O['начало'].map(_t), конец=O['конец'].map(_t), доля_5xx_у_страниц_в_пик=(O['доля_5xx_у_страниц_в_пик'] * 100).round(0).astype(int))
        put('Сбои', O, {'сбой': 'Сбой', 'начало': 'Начало', 'конец': 'Конец', 'минут_всего': 'Длительность (мин)', 'пик_минут': 'Острая фаза (мин)', 'запросов': 'Запросов',
             'запросов_в_мин': 'Запросов в минуту', 'обычно_запросов_в_мин': 'В норме', 'обрывов_499': 'Обрывов 499', 'страниц_с_5xx': 'Страниц с 5xx',
             'доля_5xx_у_страниц_в_пик': 'Доля 5xx на пике, %', 'вероятная_причина': 'Картина сбоя'},
            {'Сбой': 13, 'Начало': 16, 'Конец': 16, 'Картина сбоя': 60}, wrap=('Картина сбоя',), kpi=[('Сбоев', len(O)), ('Минут', int(O['минут_всего'].sum()))], kpi_col='В норме')
        notes_of(wb['Сбои'])
    B = S.get('Битые адреса из скриптов')
    if B is not None and len(B):
        act = list(B['актуально']) if 'актуально' in B else [str(v) >= recent for v in B.get('Последний', [''] * len(B))]   # признак — из движка (derive)
        B = B.drop(columns=[c_ for c_ in ('актуально',) if c_ in B])
        B = B.assign(Последний=B['Последний'].map(_d)) if 'Последний' in B else B
        if 'Коды' in B: B = B.assign(Коды=B['Коды'].map(fmt_codes)).rename(columns={'Коды': 'Ошибки'})   # «404 (7)» — как везде, ошибки красным
        data_sheet(wb, 'Битые из скриптов', B, {'Адрес': 50, 'Страница-источник': 40}, wrap=('Адрес', 'Страница-источник'), red=(),
                   row_rule=lambda r, _it=iter(act): F_NOTE if next(_it) else None, font_rule=lambda r, _it=iter(act): None if next(_it) else T_GREY)
        notes_of(wb['Битые из скриптов'])
        legend_of(wb['Битые из скриптов'], ['люди', 'неактуально'])
    # «Поисковые ошибки» — только в 06 «SEO» (каталог #8: запасных мест SEO-листов в 02–05 нет)
    A = S.get('Реклама: посадочные с ошибками')
    if A is not None and len(A):
        act = list(A['актуально']) if 'актуально' in A else [str(v) >= recent for v in A['последний']]   # признаки — из движка (derive)
        bad = list(A['впустую']) if 'впустую' in A else [int(c_) != 499 and a_ for c_, a_ in zip(A['entry_status'], act)]
        put('Рекламные ошибки', A.assign(первый=A['первый'].map(_d), последний=A['последний'].map(_d), entry_status=A['entry_status'].astype(int).astype(str)), {'entry': 'Посадочная', 'entry_status': 'Код', 'кликов': 'Кликов', 'первый': 'Первый', 'последний': 'Последний'}, {'Посадочная': 55}, wrap=('Посадочная',), center=('Код',),
            rule=lambda r, _it=iter(bad): F_NOTE if next(_it) else None, font=lambda r, _it=iter(act): None if next(_it) else T_GREY)
        from openpyxl.styles import Font as F_
        from .report_index import RED as R_
        ws_ = wb['Рекламные ошибки']
        kc = next((c.column for c in ws_[4] if c.value == 'Код'), None)
        for r_ in range(5, 5 + len(A)):   # 404 и 5xx — красным; 499 (не дождался) — нет
            c_ = ws_.cell(r_, kc) if kc else None
            if c_ is not None and str(c_.value) != '499': c_.font = F_(name='Arial', size=9, color=R_)
        notes_of(wb['Рекламные ошибки'])
        legend_of(wb['Рекламные ошибки'], ['клик_впустую', 'обычно', 'неактуально'])
    soft_sheet(wb, S.get('Идентичные ответы'), S.get('Пустые страницы'))
    L = S.get('Error-лог')
    if L is not None and len(L):
        import re as re_
        L = L.assign(первый=L['первый'].map(_d), последний=L['последний'].map(_d), значит=L['тип'].map(lambda t: ERRLOG_MEANING.get(t, ('', False))[0]),
                     запрос=L['запрос'].map(lambda q: '\n'.join(f"{a.strip()} ({n})" for a, n in re_.findall(r'(.+?):\s*(\d+)(?:,\s*|$)', str(q))) or str(q)))
        norm = [ERRLOG_MEANING.get(t, ('', False))[1] for t in L['тип']]
        put('Error-лог', L, {'тип': 'Тип', 'значит': 'Диагностическое значение', 'сообщений': 'Сообщений', 'первый': 'Первый', 'последний': 'Последний', 'запрос': 'Запросы-инициаторы', 'пример': 'Типовое сообщение'},
            {'Тип': 26, 'Диагностическое значение': 40, 'Запросы-инициаторы': 55, 'Типовое сообщение': 55}, wrap=('Тип', 'Диагностическое значение', 'Запросы-инициаторы', 'Типовое сообщение'),
            rule=lambda r, _it=iter(norm): F_GREY if next(_it) else None)
        notes_of(wb['Error-лог'])
        legend_of(wb['Error-лог'], ['норма', 'обычно'])


def more(W, code, shown, total, names, what=''):
    """Под блоком сводки: сколько показано из скольких — и ссылка на полный лист (каталог: ссылка блока code с текстом auto «more»)."""
    txt = f'Показаны {shown} из {total}' if total > shown else 'Подробно'
    if what and total > shown: txt = f'{txt} {what}'
    W.links_of(code, names, more=txt)


def overview(wb, res, S, site):
    """Первый лист файла: шапка с логотипом и сводка ошибок блоками, у каждого — подпись и ссылка на полный лист."""
    E = res.get('errors') or {}
    sm = E.get('сводка') or {}
    if 'Обзор' in wb.sheetnames: del wb['Обзор']
    ws = wb.create_sheet('Обзор', 0)
    W = Wide(ws)
    brand_header(ws, W, res, last='I', site=site.upper())
    names = set(wb.sheetnames)
    K = sm.get('коды')
    if K is not None and len(K):
        W.block('responses')
        rows = [[r['кто'], int(r['запросов']), int(r['ок']), int(r['e4']), int(r['e5']), int(r['n499']),
                 f"{(r['e4'] + r['e5']) / max(1, r['запросов']) * 100:.1f}%".replace('.', ',')] for _, r in K.iterrows() if r['запросов']]
        W.table(['Кто', 'Запросов', 'Без ошибок', 'Ошибки 4xx', 'Ошибки 5xx', '499 — не дождались', 'Доля ошибок'], rows, ['B', 'C', 'D', 'E', 'F', 'G', 'HI'], num=(1, 2, 3, 4, 5, 6))
        W.links_of('responses', names)
    X = S.get('Сбои')
    if X is not None and len(X):
        W.block('outages')
        O = E.get('сбои') if E.get('сбои') is not None and len(E.get('сбои')) else X.assign(сбой='', запросов=0)
        rows = [[r['сбой'], _t(r['конец']), int(r['минут_всего']), int(r['запросов']), int(r['обрывов_499']), int(r['страниц_с_5xx'])] for _, r in O.iterrows()]
        W.table(['Сбой', 'Конец', 'Длительность (мин)', 'Запросов', 'Обрывов 499', 'Страниц с 5xx'], rows, ['B', 'C', 'D', 'EF', 'G', 'HI'], num=(2, 3, 4, 5), wrap=0.95)   # картина сбоя — на листе «Сбои»
        W.links_of('outages', names)
    P = sm.get('страницы')
    if P is not None and len(P):
        W.block('victims')
        rows = [[r['адрес'], r['сейчас'], r['почему'], int(r['запросов']), int(r['Люди'])] for _, r in P.iterrows()]
        W.table(['Страница', 'Текущий статус', 'Основание', 'Запросов', 'Люди'], rows, ['BC', 'DE', 'FG', 'H', 'I'], num=(3, 4), wrap=0.95)   # в сводке всё важное — без заливки
        more(W, 'victims', len(P), len(E.get('нерабочие', [])), names)
    G = sm.get('разделы')
    if G is not None and len(G):
        W.block('sections')
        rows = [[r['раздел'], int(r['адресов']), int(r['Люди']), int(r['Поисковики']), int(r['Свои'])] for _, r in G.head(10).iterrows()]
        W.table(['Раздел', 'Страниц', 'Люди', 'Поисковики', 'Свои'], rows, ['BC', 'D', 'E', 'F', 'G'], num=(1, 2, 3, 4))
    Fe = res.get('files_errors') or []
    if Fe:
        W.block('files')
        fg = (res.get('сводки') or {}).get('ошибки_файлов_по_группам')
        if fg is None: fg = pd.DataFrame(Fe).groupby('группа').agg(файлов=('файлов', 'sum'), запросов=('запросов', 'sum'), примеры=('адрес', lambda s: list(s)[:2])).sort_values('запросов', ascending=False)
        rows = [[g, int(r['файлов']), int(r['запросов']), '\n'.join(map(str, r['примеры']))] for g, r in fg.iterrows()]
        W.table(['Группа', 'Файлов', 'Запросов', 'Например'], rows, ['B', 'C', 'D', 'EFGHI'], num=(1, 2), wrap=0.95)
        W.links_of('files', names)
    BI = E.get('битые_внутренние')
    if BI is not None and len(BI):
        W.block('broken_internal')
        rows = [[r['откуда'], r['куда'], int(r['переходов']), int(r['людей'])] for _, r in BI.head(5).iterrows()]
        W.table(['Где стоит ссылка', 'Куда ведёт', 'Переходов', 'Люди'], rows, ['BC', 'DEF', 'G', 'H'], num=(2, 3), wrap=0.95)
        more(W, 'broken_internal', min(5, len(BI)), len(BI), names)
    BX = E.get('битые_внешние')
    if BX is not None and len(BX):
        W.block('broken_external')
        rows = [[r['вид'], r['откуда'], r['куда'], int(r['переходов']), int(r['людей'])] for _, r in BX.head(5).iterrows()]
        W.table(['Откуда', 'Сайт', 'Куда ведёт', 'Переходов', 'Люди'], rows, ['B', 'C', 'DEF', 'G', 'H'], num=(3, 4), wrap=0.95)
        more(W, 'broken_external', min(5, len(BX)), len(BX), names)
    BJ = S.get('Битые адреса из скриптов')
    if BJ is not None and len(BJ):
        W.block('broken_scripts')
        rows = [[r['Адрес'], int(r['Запросов со страниц сайта']), r['Страница-источник']] for _, r in BJ.head(5).iterrows()]
        W.table(['Адрес', 'Запросов', 'Страница'], rows, ['BCD', 'E', 'FGHI'], num=(1,))
        more(W, 'broken_scripts', min(5, len(BJ)), len(BJ), names)
    SE = S.get('Ошибки у поисковиков')
    if SE is not None and len(SE):
        W.block('search_errors')
        g = (res.get('сводки') or {}).get('поиск_по_роботам')
        if g is None: g = SE.groupby('робот').agg(типов=('шаблон', 'nunique'), запросов=('запросов', 'sum'), ошибок=('ошибок', 'sum')).sort_values('ошибок', ascending=False).reset_index()
        W.table(['Робот', 'Типов страниц', 'Запросов', 'Ошибок'], [[r['робот'], int(r['типов']), int(r['запросов']), int(r['ошибок'])] for _, r in g.iterrows()], ['B', 'C', 'D', 'E'], num=(1, 2, 3))
        more(W, 'search_errors', len(g), len(SE), names, what='типов страниц по роботам')
    AD = S.get('Реклама: посадочные с ошибками')
    if AD is not None and len(AD):
        W.block('ad_errors')
        rows = [[r['entry'], int(r['entry_status']), int(r['кликов'])] for _, r in AD.head(5).iterrows()]
        W.table(['Посадочная', 'Код', 'Кликов'], rows, ['BCD', 'E', 'F'], num=(2,), center=(1,))
        more(W, 'ad_errors', min(5, len(AD)), len(AD), names)
    if 'Мягкие ошибки' in names:
        W.block('soft')
        W.links_of('soft', names)
    EL = S.get('Error-лог')
    if EL is not None and len(EL):
        W.block('error_log')
        rows = [[r['тип'], ERRLOG_MEANING.get(r['тип'], ('', False))[0], int(r['сообщений']), _d(r['первый']), _d(r['последний'])] for _, r in EL.iterrows()]
        W.table(['Тип', 'Диагностическое значение', 'Сообщений', 'Первый', 'Последний'], rows, ['B', 'CDE', 'F', 'G', 'H'], num=(2,), center=(3, 4), wrap=0.95)
        W.links_of('error_log', names)
    W.related(res, 'Ошибки')
    W.finish()   # блоки — в порядке каталога (#11)
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def slow_sheet(wb, T):
    """«Признаки торможения»: когда сервер не успевает — обрывы 499 у людей, ошибки шлюза, таймауты, время ответа (если есть в логе)."""
    if not T: return None
    from .derive import slowness_text
    intro = T.get('вывод') or slowness_text(T)   # вывод — из движка (derive)
    blocks = [('drops', T.get('499'), '0.0'), ('gateway', T.get('шлюз'))]
    if T.get('таймауты') is not None: blocks.append(('timeouts', T['таймауты']))
    if T.get('время') is not None: blocks.append(('response_time', T['время'], '0.00'))
    return heat_sheet(wb, 'Признаки торможения', blocks, unit='', conclusion=intro)


def build(wb, res, S, site):
    E = res.get('errors') or {}
    T, Sr = E.get('коды'), E.get('источники')
    if 'Статистика' not in wb.sheetnames: wb.create_sheet('Статистика')
    def kpi(body):
        h = body['5xx|Люди'] + body['4xx|Люди']
        return [('Ошибок у людей в день', int(round(h[body['_полный']].mean())) if body['_полный'].any() else int(round(h.mean()))),
                ('Больше всего', _day(body.loc[h.idxmax(), 'день']) if len(body) else ''), ('5xx у людей', int(body['5xx|Люди'].sum()))]
    journal_sheet(wb, 'Статистика', E.get('журнал'), kpi, kpi_col='4xx|Люди', outages=E.get('сбои'))
    slow_sheet(wb, E.get('торможение'))
    for fam, _ in FAMILIES:
        if fam not in wb.sheetnames: wb.create_sheet(fam)
        if code_tab(wb, T, Sr, fam) is None: del wb[fam]
    victims(wb, E)
    if res.get('files_errors'):
        service_files(wb, dict(res, files=res['files_errors']), 'Ошибки файлов', errors_mode=E.get('с_дня', True))
    links_sheets(wb, E)
    full_sheets(wb, S, E.get('с_дня', ''), E.get('сбои'))
    from .catalog import prune
    prune(wb, 'Ошибки')   # до Обзора: строки на невыведенные листы не появляются (#8)
    overview(wb, res, S, site)
    redden_codes(wb, only=('Битые внутренние', 'Битые внешние', 'Поисковые ошибки', 'Пострадавшие страницы', '5xx', '403 и прочие 4xx', 'Битые из скриптов'), rx=r'^Ошибки$')   # красный — где он различает 404 и 5xx
    byname = {w.title: w for w in wb._sheets}
    head_ = [byname[n_] for n_ in catalog_order('Ошибки') if n_ in byname]   # порядок — из каталога (#8)
    wb._sheets = head_ + [w for w in wb._sheets if w not in head_]
