"""NXLD: простые листы-таблицы в общем стиле (заголовок, пояснение, оранжевая шапка, фильтр, закреплённая шапка)."""
import numbers
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, PatternFill, Side
from .report_index import ORANGE, ORANGE2, INK, GREY, F_NOTE, NUM_FMT, RED

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as ILLEGAL
SEP = Side(style='thin', color='D9D9D9')
WHITE = Side(style='thin', color='FFFFFF')


def data_sheet(wb, name, df, title, note='', widths=None, wrap=(), fill_rule=None, bold_rule=None, center=(), sort_by=None, kpi=None, kpi_col=None, links=(), size=9, row_rule=None, red=('Ошибки',), font_rule=None):
    """Пересобирает лист name: строка 1 — заголовок, 2 — пояснение, 4 — шапка, дальше данные.
    widths — ширины колонок; wrap — колонки с переносом; fill_rule(col, value) → цвет заливки или None."""
    idx = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name, idx)
    ws.sheet_view.showGridLines = True
    cols = list(df.columns)
    ws.column_dimensions['A'].width = 2.5          # поле слева, как на остальных листах
    ws['B1'] = title.upper(); ws['B1'].font = Font(name='Montserrat', size=16, bold=True, color=ORANGE)
    ws.row_dimensions[1].height = 34
    if note:   # подзаголовок — фирменным Comfortaa, как на обзоре
        ws['B2'] = note; ws['B2'].font = Font(name='Comfortaa', size=11, bold=True, color=GREY)
        ws['B2'].alignment = Alignment(vertical='center', wrap_text=False)
    ws.row_dimensions[2].height = 24
    X = 1   # сдвиг колонок
    if kpi:   # цифры сводки — справа от заголовка: подпись мелко, число крупно
        c0 = (cols.index(kpi_col) if kpi_col in cols else max(0, len(cols) - len(kpi))) + 1 + X
        for k, (lab, val) in enumerate(kpi):
            a = ws.cell(1, c0 + k, lab); a.font = Font(name='Arial', size=9, color=INK); a.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            b = ws.cell(2, c0 + k, val); b.font = Font(name='Arial', size=14, bold=True, color='000000'); b.alignment = Alignment(horizontal='center', vertical='center')
            if isinstance(val, int) and val >= 1000: b.number_format = NUM_FMT
    H = 4
    for j, c in enumerate(cols, 1):
        cell = ws.cell(H, j + X, c)
        cell.font = Font(name='Arial', size=size, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=ORANGE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=WHITE, right=WHITE)
    ws.cell(H, 1 + X)
    ws.row_dimensions[H].height = 32
    for i, row in enumerate(df.itertuples(index=False), H + 1):
        rf = row_rule(dict(zip(cols, row))) if row_rule else None   # заливка всей строки
        ff = font_rule(dict(zip(cols, row))) if font_rule else None   # цвет текста всей строки (неактуальное — серым)
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v): v = None
            if hasattr(v, 'item'): v = v.item()
            if isinstance(v, str): v = ILLEGAL.sub('�', v)   # управляющие символы из адресов сканеров Excel не принимает
            cell = ws.cell(i, j + X, v)
            col = cols[j - 1]
            isnum = isinstance(v, numbers.Number) and not isinstance(v, bool)
            cell.font = Font(name='Arial', size=size, bold=bool(bold_rule and bold_rule(col, v)), color=ff or (RED if col in red else '000000'))
            hz = 'center' if col in center else ('right' if isnum else 'left')
            cell.alignment = Alignment(horizontal=hz, vertical='top', wrap_text=col in wrap, indent=1 if hz != 'center' else 0)
            if isnum and isinstance(v, int) and abs(v) >= 1000: cell.number_format = NUM_FMT
            cell.border = Border(bottom=SEP)
            f = rf or (fill_rule(col, v) if fill_rule else None)
            if f: cell.fill = PatternFill('solid', fgColor=f)
    for j, c in enumerate(cols, 1):
        letter = ws.cell(H, j + X).column_letter
        ws.column_dimensions[letter].width = (widths or {}).get(c) or min(45, max(10, len(str(c)) + 2, int(df[c].astype(str).str.len().quantile(0.9)) + 2 if len(df) else 10))
    r_ = H + len(df) + 2   # ссылки — под таблицей, как на других листах
    for text, target in links:
        if isinstance(target, tuple):   # лист другого файла отчёта
            c = ws.cell(r_, 1 + X, f'{text}: {target[0]}, лист «{target[1]}» →'); c.hyperlink = f"{target[0]}#'{target[1]}'!A1"
        else:
            c = ws.cell(r_, 1 + X, f'{text}: лист «{target}» →'); c.hyperlink = f"#'{target}'!A1"
        c.font = Font(name='Arial', size=10, color=ORANGE2, underline='single'); r_ += 1
    ws.freeze_panes = ws.cell(H + 1, 1 + X)
    if len(df):
        ws.auto_filter.ref = f"{ws.cell(H, 1 + X).coordinate}:{ws.cell(H + len(df), len(cols) + X).coordinate}"
        if sort_by and sort_by in cols:   # отметка сортировки в фильтре — Excel покажет, по чему отсортировано
            c_ = cols.index(sort_by) + 1 + X
            ws.auto_filter.add_sort_condition(f"{ws.cell(H + 1, c_).coordinate}:{ws.cell(H + len(df), c_).coordinate}")
    return ws


LEGEND = {'критично': (F_NOTE, None, 'Критично — см. «Почему важно»'), 'обычно': ('FFFFFF', None, 'Ошибка есть, не критична'),
          'неактуально': ('FFFFFF', GREY, 'Уже неактуально: починилось или давно не встречалось'),
          'чужое': ('EFEFEF', None, 'Чужое: сканеры и посторонние'), 'норма': ('EFEFEF', None, 'Норма'),
          'люди': (F_NOTE, None, 'Ошибку получили люди, и она встречается сейчас'),
          'со_страниц': (F_NOTE, None, 'Файл просят страницы сайта, и ошибка встречается сейчас'),
          '5xx': (F_NOTE, None, 'Сервер отвечает ошибкой 5xx'), 'клик_впустую': (F_NOTE, None, 'Клик привёл на ошибку: деньги впустую'),
          'сбой': (F_NOTE, None, 'Сбой: сервер отвечал ошибками на многие страницы сразу'), 'выходные': ('F3F3F3', None, 'Выходные'),
          'неполный': ('FFFFFF', GREY, 'Неполный день: в лог попала только часть суток'),
          'находка': (F_NOTE, None, 'Сервер отдал постороннему то, что ему не положено'),
          'не_улика': ('EFEFEF', None, 'Только обычные страницы сайта и заглушки — не улика'),
          'чужой': (F_NOTE, None, 'Чужой: не сотрудник'), 'программа': ('FFFFFF', GREY, 'Не человек: программа или бот'),
          'ложный': ('EFEFEF', None, 'Ложный раздел: на сайте его нет, сканеры угадывали'),
          'открыто': (F_NOTE, None, 'Настоящий раздел: есть страницы, которые открываются без входа'),
          'всплеск': (F_NOTE, None, 'Всплеск от программ, похожий на DDoS, или с последствиями для людей'),
          'подозрительно': (F_NOTE, None, 'Сервер ответил 200, и ответ не похож на обычную страницу — проверить'),
          'посторонние': (F_NOTE, None, 'Адрес с токеном видели посторонние'), 'свои_норма': ('EFEFEF', None, 'Видели только свои — норма'),
          'ловушка': (F_NOTE, None, 'По вариантам ходят в основном роботы и боты'),
          'принято': (F_NOTE, None, 'Сервер принял необычный метод (ответ 2xx)'),
          'утечка': (F_NOTE, None, 'Файл отдаётся посторонним сейчас или не проверен'), 'закрыт': ('FFFFFF', GREY, 'Закрыт: больше не отдаётся'),
          'массовые_люди': ('FFFFFF', GREY, 'Люди и свои: проверить, не общий ли это IP (офис, мобильный оператор)')}


def legend(ws, keys):
    """Легенда цветов под таблицей: квадратик цвета и подпись. keys — какие цвета есть на этом листе (ключи LEGEND)."""
    r_ = ws.max_row + 2
    for k in keys:
        fill, font, text = LEGEND[k]
        c0 = ws.cell(r_, 2, text)   # подпись прямо на образце цвета: залитая ячейка с текстом нужного цвета
        c0.fill = PatternFill('solid', fgColor=fill); c0.border = Border(left=SEP, right=SEP, top=SEP, bottom=SEP)
        c0.font = Font(name='Arial', size=9, color=font or '000000')
        c0.alignment = Alignment(vertical='center', indent=1)
        r_ += 1


def conversions(wb, C, name='Конверсии'):
    """Все отправки форм: кто, какая форма, принята ли, откуда пришёл."""
    if C is None or not len(C) or name not in wb.sheetnames: return
    from .anatomy import name_of, FORM_NAME
    import re
    d = pd.DataFrame({
        'Время': pd.to_datetime(C['время']).dt.strftime('%d.%m.%Y %H:%M'),
        '_t': pd.to_datetime(C['время']),
        'Кто': C['группа'].astype(str),
        'Форма': C['цель'].astype(str).map(lambda g: (name_of(g, FORM_NAME) or ('общий обработчик форм' if re.search(r'/form\.php$', g) else 'форма')).capitalize()),
        'Принята': C['принята'].astype(str),
        'Статус': C['статус'].astype(str) if 'статус' in C else '',
        'Код ответа': pd.to_numeric(C['код'], errors='coerce').astype('Int64'),
        'Канал': C['канал'].astype(str),
        'Страница входа': C['вход'].astype(str),
        'Откуда пришёл': C['вход_реферер'].astype(str).replace({'-': '', 'nan': ''}),
        'Страниц до отправки': pd.to_numeric(C['страниц_до'], errors='coerce').astype('Int64'),
        'Секунд от входа': pd.to_numeric(C['сек_от_входа'], errors='coerce').astype('Int64'),
        'Почему бот': C['подгруппа'].astype(str).where(C['группа'] == 'Боты', '').str.replace('спам форм: ', '', regex=False),
        'IP': C['ip'].astype(str),
        'Сеть': C['сеть'].astype(str),
        'Обработчик': C['цель'].astype(str),
    }).sort_values('_t').drop(columns='_t')   # по времени
    d = d.astype(object).where(d.notna(), None)
    n_ = lambda g: int((C['группа'] == g).sum())
    kpi = [('Отправок', len(C)), ('Люди', n_('Люди')), ('Принято у людей', int(((C['группа'] == 'Люди') & (C['принята'] == 'да')).sum())), ('Боты', n_('Боты')), ('Свои', n_('Свои'))]
    widths = {'Время': 17, 'Кто': 13, 'Форма': 22, 'Принята': 10, 'Статус': 44, 'Код ответа': 9, 'Канал': 18, 'Страница входа': 28, 'Откуда пришёл': 24,
              'Страниц до отправки': 11, 'Секунд от входа': 10, 'Почему бот': 34, 'IP': 16, 'Сеть': 22, 'Обработчик': 50}
    fill = None
    row_rule = lambda r: F_NOTE if r.get('Кто') == 'Боты' else ('EFEFEF' if r.get('Кто') == 'Свои' else None)
    bold = lambda col, v: col == 'Принята' and v == 'да'
    data_sheet(wb, name, d, 'Конверсии', 'Все отправки форм за период', widths, wrap=('Почему бот', 'Статус'), fill_rule=fill, bold_rule=bold,
               center=('IP', 'Код ответа', 'Принята'), kpi=kpi, kpi_col='Канал', links=[('Сводка по формам', 'Анатомия сайта')], row_rule=row_rule)


# ---- точки приёма данных и POST-отправки ----
SITE_KINDS = ('Заявка', 'Заявка?', 'Вход', 'Обмен с 1С', 'API', 'Вебхук', 'Фильтр каталога', 'Поиск по сайту', 'Форма (GET)', 'Пагинация', 'Тип отображения', 'Сортировка', 'Служебный скрипт', 'Подгрузка на странице', 'Админка', 'Загрузка файлов')


def _codes(s):
    import re
    return {int(k): int(v) for k, v in re.findall(r'(\d{3}):\s*(\d+)', str(s))}


def classify_post(r, login_roots, engine, ev=None, prof=None):
    """Что это за адрес и почему — человеческим языком."""
    import re
    a, out = str(r['адрес']), str(r.get('вывод', ''))
    cd = _codes(r.get('коды'))
    tot = max(1, sum(cd.values()))
    ok = sum(v for k, v in cd.items() if 200 <= k < 400 and k not in (301,)) / tot
    if out.startswith('цель'):
        e = (ev or {}).get(a)
        if e and e['редиректов']:
            how = ', '.join(f"{k} — {v}" for k, v in sorted(e['как'].items(), key=lambda x: -x[1]))
            if e['подтверждено'] == e['редиректов']:
                return 'Заявка', f"после отправки — переадресация (302), затем {how}: подтверждено {e['подтверждено']} из {e['редиректов']}"
            if e['подтверждено']:
                return 'Заявка', f"после отправки — переадресация (302), затем {how}: подтверждено {e['подтверждено']} из {e['редиректов']}; остальные не приняты"
            return 'Заявка', f"переадресация (302) есть, но подтверждения успеха нет ни в одном из {e['редиректов']} случаев — заявки не приняты"
        return 'Заявка', ('переадресация после отправки (3xx)' if '3xx' in out else 'успех по коду ответа не виден (всегда 200) — не подтверждено')
    if any(a == root or a.startswith(root) and a.rstrip('/') == root.rstrip('/') for root in login_roots) or (a in login_roots):
        return 'Вход', 'форма входа: неудачный вход возвращает ту же страницу, удачный — другую'
    if re.search(r'/bitrix/admin/|/wp-admin/|/administrator/', a):
        return 'Админка', 'запросы из админки движка (работа сотрудников)'
    if re.search(r'/wp-|wordpress|xmlrpc|^/wp/', a) and engine != 'WordPress':
        return 'Сканер', 'адрес WordPress, а сайт на другом движке'
    bad = cd.get(404, 0) + cd.get(405, 0) + cd.get(301, 0) + cd.get(302, 0)
    if cd and cd.get(404, 0) + cd.get(405, 0) >= 0.8 * tot:
        return 'Сканер', 'такой страницы нет на сайте (404)'
    if cd and bad >= 0.8 * tot:
        return 'Сканер', 'страницы нет (404) или перенаправление (301) — форму не принимает'
    if cd and cd.get(403, 0) >= 0.5 * tot:
        return 'Сканер', 'защита отказала (403)'
    if 'upload' in a:
        return 'Загрузка файлов', 'загрузка файлов на сервер'
    if any(a.startswith(root) for root in login_roots):
        return 'Служебный скрипт', 'скрипт внутри закрытого раздела'
    if re.search(r'ajax|/tools/|/services/|autosave|\.php$', a) and a not in ('/index.php',) and ok >= 0.5:
        return 'Служебный скрипт', 'скрипт сайта: подгружает данные, не заявка'
    p = (prof or {}).get(a, {})
    if p.get('свои', 0) >= 0.8:
        return 'Подгрузка на странице', f"POST шлют браузеры посетителей, которые уже на сайте ({int(p['свои'] * 100)}%) — страница подгружает данные (список, форму)"
    if p:
        tail = f", например параметр «{p['параметр']}»" if p.get('параметр') else ''
        return 'Сканер', f"POST на обычную страницу без перехода с сайта ({int((1 - p.get('свои', 0)) * 100)}% запросов){tail} — боты и сканеры"
    if ok >= 0.5:
        return 'Подгрузка на странице', 'обычная страница отвечает на POST'
    return 'Сканер', 'обычная страница, форму не принимает — отправляют боты и сканеры'


def _post_rows(res):
    m = res.get('site_map') or {}
    F = m.get('forms') or []
    if isinstance(F, str):
        try: F = eval(F)
        except Exception: F = []
    OS = res.get('sheets', {}).get('Нагрузка и безопасность', {}).get('Открытые служебные разделы', pd.DataFrame())
    roots = set(OS.loc[OS['форма_входа'] == 'да', 'раздел'].astype(str)) if len(OS) and 'форма_входа' in OS else set()
    engine = ((m.get('engines') or [{}])[0] or {}).get('движок', '')
    rows = []
    for f in F:
        if not isinstance(f, dict): continue
        kind, why = classify_post(f, roots, engine, res.get('form_evidence'), (res.get('anatomy') or {}).get('post_pages'))
        rows.append({'Адрес точки': f['адрес'], 'Метод': 'POST', 'Опознано как': kind, 'Запросов': int(f.get('отправок') or 0), 'Уникальных IP': int(f.get('IP') or 0), 'Улики': why,
                     'Коды ответа': str(f.get('коды', '')), '_t0': pd.to_datetime(f.get('первый')), '_t1': pd.to_datetime(f.get('последний'))})
    return rows


def _finish(rows):
    d = pd.DataFrame(rows)
    if not len(d): return d
    d['Первый'] = d['_t0'].dt.strftime('%d.%m.%Y %H:%M'); d['Последний'] = d['_t1'].dt.strftime('%d.%m.%Y %H:%M')
    sp_ = d['Коды ответа'].map(split_codes)
    d['Ответы'], d['Ошибки'] = sp_.str[0], sp_.str[1]
    d = d.sort_values('Запросов', ascending=False).drop(columns=['_t0', '_t1', 'Коды ответа'])   # по числу запросов; группы — фильтром
    order = ['Адрес точки', 'Запросов', 'Уникальных IP', 'Метод', 'Опознано как', 'Улики', 'Ответы', 'Ошибки', 'Первый', 'Последний']
    d = d[[c_ for c_ in order if c_ in d.columns]]
    return d


WIDTHS = {'Адрес точки': 55, 'Метод': 8, 'Опознано как': 22, 'Запросов': 10, 'Уникальных IP': 11, 'Улики': 50, 'Ответы': 30, 'Ошибки': 24, 'Первый': 16, 'Последний': 16}


def intake(wb, res, name='Точки приёма данных', before='Конверсии'):
    """Overview: только то, что сайт реально принимает — заявки, вход, фильтры и поиск, свои скрипты, админка."""
    rows = [r for r in _post_rows(res) if r['Опознано как'] != 'Сканер']
    for g in (res.get('anatomy') or {}).get('api', []) + (res.get('anatomy') or {}).get('get_приём', []):
        rows.append({'Адрес точки': g['адрес'], 'Метод': 'GET', 'Опознано как': g['что'], 'Запросов': g['отправок'], 'Уникальных IP': g['IP'], 'Улики': g['почему'],
                     'Коды ответа': g['коды'], '_t0': pd.to_datetime(g['первый']), '_t1': pd.to_datetime(g['последний'])})
    d = _finish(rows)
    if not len(d): return
    if name not in wb.sheetnames: wb.create_sheet(name, wb.sheetnames.index(before) if before in wb.sheetnames else len(wb.sheetnames))
    cnt = d['Опознано как'].value_counts()
    cat_ = int(sum(cnt.get(k, 0) for k in ('Фильтр каталога', 'Поиск по сайту', 'Форма (GET)', 'Пагинация', 'Тип отображения', 'Сортировка')))
    serv_ = int(sum(cnt.get(k, 0) for k in ('Служебный скрипт', 'Подгрузка на странице', 'Админка', 'Загрузка файлов')))
    api_ = int(sum(cnt.get(k, 0) for k in ('Обмен с 1С', 'API', 'Вебхук')))
    kpi = [('Точек приёма', len(d)), ('Заявки', int(cnt.get('Заявка', 0))), ('Поиск и навигация', cat_),
           ('Вход', int(cnt.get('Вход', 0))), ('Служебные', serv_), ('API и обмен', api_)]
    kpi = kpi[:1] + sorted(kpi[1:], key=lambda x: x[1] == 0)   # нули — в конец
    row_rule = lambda r: 'EFEFEF' if r.get('Опознано как') in ('Служебный скрипт', 'Подгрузка на странице', 'Админка', 'Загрузка файлов') else None
    bold = lambda col, v: col == 'Опознано как' and v in ('Заявка', 'Вход')
    data_sheet(wb, name, d, 'Точки приёма данных', 'Где сайт принимает данные', WIDTHS,
               wrap=('Улики', 'Адрес точки', 'Ответы', 'Ошибки'), bold_rule=bold, center=('Метод',), kpi=kpi, kpi_col='Метод', row_rule=row_rule,
               links=[('Все отправки форм', 'Конверсии'), ('Сводка по формам', 'Анатомия сайта')])
    ws = wb[name]   # сноска и чего нет — тоже результат
    if d['Адрес точки'].astype(str).str.contains('фасетн').any():
        r0 = ws.max_row + 2
        c = ws.cell(r0, 2, 'Основные страницы — адреса структуры сайта. Фасетные страницы — адреса, которые генерирует фильтр каталога из комбинаций условий (…/filter/…/apply/); в структуре сайта их нет.')
        c.font = Font(name='Arial', size=10, italic=True, color=INK)
    absent = [lab for lab, keys in (('Поиск по сайту', ('Поиск по сайту',)), ('API и обмен с 1С, CRM, вебхуки', ('Обмен с 1С', 'API', 'Вебхук'))) if not any(cnt.get(k, 0) for k in keys)]
    if absent:
        r_ = ws.max_row + 2
        c = ws.cell(r_, 2, 'Обращений не найдено: ' + '; '.join(absent) + '.')
        c.font = Font(name='Arial', size=10, italic=True, color=INK)


def post_all(wb, res, name='POST-отправки', after='GET-отправки'):
    """03 Безопасность: все POST-отправки, включая сканеры."""
    d = _finish(_post_rows(res))
    if not len(d): return
    d = d.drop(columns=['Метод'])
    if name not in wb.sheetnames:
        pos = wb.sheetnames.index(after) + 1 if after in wb.sheetnames else len(wb.sheetnames)
        wb.create_sheet(name, pos)
    sc = d[d['Опознано как'] == 'Сканер']
    kpi = [('Адресов', len(d)), ('Принимает сайт', len(d) - len(sc)), ('Адресов сканеров', len(sc)), ('Запросов сканеров', int(sc['Запросов'].sum()))]
    row_rule = lambda r: F_NOTE if r.get('Опознано как') == 'Сканер' else None
    data_sheet(wb, name, d, name, 'Все адреса, куда за период отправляли данные методом POST, — и сайт, и сканеры', WIDTHS,
               wrap=('Улики', 'Адрес точки', 'Ответы', 'Ошибки'), center=(), kpi=kpi, kpi_col='Метод', row_rule=row_rule)
    ws = wb[name]   # ссылка на другой файл
    r_ = ws.max_row + 2
    c = ws.cell(r_, 2, 'Что сайт принимает на самом деле: лист «Точки приёма данных» в NXLD_01_Overview.xlsx →')
    c.hyperlink = "NXLD_01_Overview.xlsx#'Точки приёма данных'!A1"; c.font = Font(name='Arial', size=10, color=ORANGE2, underline='single')


def split_codes(s):
    """«200:65634, 404:12, 499:3» → («200 (65 634), 499 (3)», «404 (12)»): ответы и ошибки (4xx/5xx, кроме 499 — человек ушёл сам)."""
    ok, bad = [], []
    for k, v in _codes(s).items():
        (bad if k >= 400 and k != 499 else ok).append(f"{k} ({int(v):,})".replace(',', '\u00a0'))
    return ', '.join(ok), ', '.join(bad)


def fmt_codes(s):
    """«200:65634, 499:632» → «200 (65 634), 499 (632)» — чтобы не читалось как порт."""
    out = []
    for k, v in _codes(s).items():
        out.append(f"{k} ({int(v):,})".replace(',', '\u00a0'))
    return ', '.join(out) or str(s)



def facets_sheet(wb, F, name='Фасеты'):
    if F is None or not len(F) or name not in wb.sheetnames: return
    d = pd.DataFrame({'Раздел': F['раздел'], 'Ключ': F['ключ'], 'Значение': F['значение'], 'Условие': F['условие'], 'Расшифровка': F['расшифровка'],
                      'Запросов': F['запросов'].astype(int), 'Уникальных IP (люди)': F['людей'].astype(int), 'Откуда': F['откуда']})
    kpi = [('Условий', int(d['Условие'].nunique())), ('Значений', len(d)), ('Запросов', int(d['Запросов'].sum()))]
    widths = {'Раздел': 14, 'Ключ': 26, 'Значение': 34, 'Условие': 16, 'Расшифровка': 30, 'Запросов': 11, 'Уникальных IP (люди)': 10, 'Откуда': 30}
    data_sheet(wb, name, d, 'Фасеты', 'Что люди выбирают в фильтре каталога', widths, wrap=('Ключ', 'Значение', 'Условие', 'Расшифровка', 'Откуда'), kpi=kpi, kpi_col='Запросов',
               row_rule=lambda r: 'EFEFEF' if str(r.get('Откуда', '')).startswith('параметры') else None,
               links=[('Где фильтр принимает запросы', 'Точки приёма данных'), ('Устройство каталога', 'Анатомия сайта')])




def _who(s):
    """«браузеры/прочие:41384, Applebot:122» → «браузеры (41 384), Applebot (122)»."""
    import re
    out = []
    for part in str(s).split(', '):
        k, _, v = part.rpartition(':')
        if k and v.strip().isdigit():
            out.append(f"{k.replace('браузеры/прочие', 'браузеры')} ({int(v):,})".replace(',', '\u00a0'))
        elif part: out.append(part)
    return ', '.join(out)


def leaks_sheet(wb, L, name='Утечки служебных файлов'):
    """03: служебные файлы, которые сервер отдал посторонним, — каждый адрес проверить и закрыть."""
    if L is None or not len(L) or name not in wb.sheetnames: return None
    day = lambda v: pd.to_datetime(v).strftime('%d.%m.%Y') if v else ''
    d = pd.DataFrame({'Файл': L['файл'].astype(str), 'Отдан раз': L['ответов_200'].astype(int), 'Уникальных IP': L['IP'].astype(int),
                      'Размер ответа, байт': L['размер_у_чужих'].fillna(L['размер']).round().astype(int),
                      'Получили свои': L['получили_свои'].astype(int), 'Первый': L['первый'].map(day), 'Последний': L['последний'].map(day)})
    data_sheet(wb, name, d, name, 'Служебные файлы, которые сервер отдал посторонним', {'Файл': 50, 'Отдан раз': 11, 'Уникальных IP': 10, 'Размер ответа, байт': 14, 'Получили свои': 12, 'Первый': 12, 'Последний': 12},
               wrap=('Файл',), kpi=[('Файлов', len(d)), ('Отдан раз', int(d['Отдан раз'].sum()))], kpi_col='Уникальных IP', row_rule=lambda r: F_NOTE,
               links=[('Что искали сканеры', 'Сканеры — что искали')])
    ws = wb[name]
    r_ = ws.max_row + 2
    for t_ in ['Размер ответа не похож на страницу входа или заглушку соседних адресов — значит, отдано настоящее содержимое файла.',
               'Что сделать: открыть каждый адрес, закрыть доступ в настройках сервера (deny в nginx, правило в .htaccess) и сменить пароли и ключи, если они были в файле.']:
        c_ = ws.cell(r_, 2, t_); c_.font = Font(name='Arial', size=9, italic=True, color=GREY); r_ += 1
    return ws


def service_files(wb, res, name='Файлы', title='Файлы', note='Что забирают с сайта как отдельный файл', links=None, errors_mode=False):
    """Все файлы сайта, кроме страниц: кто забирает, откуда, коды. errors_mode — тот же лист в срезе ошибок (02 «Файлы с ошибками»):
    персиковым — файлы, которые подгружают страницы сайта, серым — те, что просят только роботы."""
    import re
    from .recon import FILE_GROUPS
    sf = res.get('files')
    if sf is None:   # старый анализ: только служебные файлы из карты
        from .anatomy import SERVICE_GROUPS
        sf = (res.get('site_map') or {}).get('service_files') or []
        if isinstance(sf, str):
            try: sf = eval(sf)
            except Exception: sf = []
        for f in sf:
            if isinstance(f, dict): f.setdefault('группа', next((nm for nm, rx in SERVICE_GROUPS if re.search(rx, str(f.get('адрес', '')), re.I)), 'Прочие'))
    rows = []
    leaked = {x['файл'] for x in res.get('leaks') or []}
    for f in sf:
        if not isinstance(f, dict): continue
        cd = _codes(f.get('коды'))
        err = sum(v for k, v in cd.items() if 400 <= k and k != 499)
        from urllib.parse import unquote
        a = unquote(str(f.get('адрес', '')), errors='replace') + (f"\n{f['внутри']}" if f.get('внутри') else '')
        if str(f.get('адрес', '')) in leaked: a += '\nотдан посторонним — см. NXLD_03, «Утечки служебных файлов»'
        day = lambda k: pd.to_datetime(f.get(k)).strftime('%d.%m.%Y') if f.get(k) else ''
        rows.append({'Файл': a, 'Размер, КБ': float(f.get('средний_размер_КБ') or 0), 'Группа': f.get('группа', 'Прочие'),
                     'Файлов': int(f.get('файлов') or 1), 'Запросов': int(f.get('запросов') or 0),
                     'Ответы': split_codes(f.get('коды'))[0], 'Ошибки': split_codes(f.get('коды'))[1], '_err': int(err),
                     **({'Браузеры': int(f.get('браузеры') or 0), 'Роботы': _who(f.get('роботы', '')) if f.get('роботы') else ''} if 'браузеры' in f
                        else {'Кто забирает': _who(f.get('кто_забирает', ''))}),
                     'Со страниц сайта': int(f.get('со_страниц') or 0), 'Напрямую': int(f.get('напрямую') or 0),
                     'Другие сайты': _who(f.get('другие_сайты', '')) if f.get('другие_сайты') else '',
                     'Первый': day('первый_день'), 'Последний': day('последний_день'), '_leak': str(f.get('адрес', '')) in leaked, '_last': str(f.get('последний_день') or '')})
    d = pd.DataFrame(rows)
    if not len(d): return None
    if 'со_страниц' not in (sf[0] if sf and isinstance(sf[0], dict) else {}) or (d['Со страниц сайта'].sum() == 0 and not d['Другие сайты'].astype(bool).any()):   # старый анализ или в логе нет Referer
        d = d.drop(columns=['Со страниц сайта', 'Напрямую', 'Другие сайты'])
    d = d.sort_values('Запросов', ascending=False)
    if name not in wb.sheetnames: wb.create_sheet(name)
    kpi = [('Файлов', int(d['Файлов'].sum())), ('Запросов', int(d['Запросов'].sum())), ('С ошибками', int((d['_err'] > 0).sum()))]
    fonts_ = [None] * len(d)
    if errors_mode:   # критично — файл просят страницы сайта, и ошибка встречается сейчас; давно не встречалась — серым текстом
        recent = errors_mode if isinstance(errors_mode, str) else ''
        own = (d['Со страниц сайта'] > 0) if 'Со страниц сайта' in d else (d['_err'] > 0)
        act = pd.Series([str(x) >= recent for x in d['_last']], index=d.index)
        fills_ = [F_NOTE if o and a else None for o, a in zip(own, act)]
        fonts_ = [None if a else GREY for a in act]
        kpi = [('Файлов', int(d['Файлов'].sum())), ('Ошибок', int(d['_err'].sum())), ('Критично', int((own & act).sum()))]
    else:
        fills_ = [F_NOTE if b_ else None for b_ in ((d['_err'] > 0) | d['_leak'])]   # ошибки и утечки — персиковым
    d = d.drop(columns=['_err', '_leak', '_last'])
    widths = {'Файл': 52, 'Размер, КБ': 10, 'Группа': 18, 'Файлов': 9, 'Запросов': 11, 'Ответы': 34, 'Ошибки': 26, 'Кто забирает': 44, 'Браузеры': 11, 'Роботы': 44, 'Со страниц сайта': 11, 'Напрямую': 11, 'Другие сайты': 40, 'Первый': 12, 'Последний': 12}
    missing = [nm for nm in ('Документы', 'Видео и звук') if nm not in set(d['Группа'])] if res.get('files') is not None and not errors_mode else []   # заметные отсутствия
    links = links if links is not None else [('Сводка по группам', 'Анатомия сайта'), ('Логи, по которым всё посчитано', 'Логи')]
    data_sheet(wb, name, d, title, note, widths, wrap=('Файл', 'Кто забирает', 'Роботы', 'Другие сайты', 'Ответы', 'Ошибки'),
               kpi=kpi, kpi_col='Файлов', links=links, row_rule=lambda r, _it=iter(fills_): next(_it), font_rule=lambda r, _it=iter(fonts_): next(_it),
               red=() if errors_mode else ('Ошибки',))
    ws = wb[name]
    if missing:
        r = ws.max_row + 2
        ws.cell(r, 2, 'Не найдено: ' + ', '.join(missing).lower() + '. Обращения сканеров к несуществующим файлам сюда не попадают.')
        ws.cell(r, 2).font = Font(name='Arial', size=9, italic=True, color=GREY)
    if errors_mode:
        r = ws.max_row + 2
        for t_ in ['Файлы, которые никто не запрашивал по ссылке, — перебор и мусор, их здесь нет: они на вкладках по кодам. Зонды сканеров — в отчёте по безопасности.']:
            ws.cell(r, 2, t_).font = Font(name='Arial', size=9, italic=True, color=GREY); r += 1
        legend(ws, ['со_страниц', 'обычно', 'неактуально'])
    return wb[name]


def _err_share(codes):
    """Доля ошибок 4xx/5xx без 499 (человек ушёл, не дождавшись) и есть ли 5xx."""
    cd = _codes(codes); tot = sum(cd.values()) or 1
    err = sum(v for k, v in cd.items() if k >= 400 and k != 499)
    return err / tot, any(k >= 500 for k in cd)


def embedded_sheet(wb, res, name='Динамические блоки'):
    """Блоки внутри страницы: что браузер подгружает сам и что человек вызывает действием."""
    E = res.get('embedded')
    if not E: return None
    def blk(e):
        n = int(e.get('блоков') or 1)
        w = 'адрес' if n % 10 == 1 and n % 100 != 11 else 'адреса' if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14 else 'адресов'
        return e['блок'] + (f' — {n} {w}' if n > 1 else '')
    day = lambda v: pd.to_datetime(v).strftime('%d.%m.%Y') if v else ''
    d = pd.DataFrame([{'Блок': blk(e), 'Размер, КБ': float(e['средний_размер_КБ']), 'Вид': e['вид'], 'Загрузка': e['загрузка'],
                       'Запросов': int(e['запросов']), 'Визитов (люди)': int(e['визитов_людей']),
                       'Автозагрузка, %': e['автозагрузка'] if e.get('автозагрузка') is not None else '',
                       'Охват страниц': int(e['охват_страниц']), 'Ответы': split_codes(e['коды'])[0], 'Ошибки': split_codes(e['коды'])[1],
                       'Первый': day(e['первый_день']), 'Последний': day(e['последний_день'])} for e in E])
    bad = [bool(split_codes(e['коды'])[1]) for e in E]
    d['_bad'] = bad
    cnt = d['Загрузка'].value_counts()
    kpi = [('Блоков', len(d)), ('Подгрузок', int(d['Запросов'].sum())), ('Авто', int(cnt.get('авто', 0))), ('По действию', int(cnt.get('по действию', 0))), ('С ошибками', int(sum(bad)))]
    if name not in wb.sheetnames: wb.create_sheet(name)
    widths = {'Блок': 52, 'Размер, КБ': 10, 'Вид': 26, 'Загрузка': 13, 'Запросов': 13, 'Визитов (люди)': 13, 'Автозагрузка, %': 13,
              'Охват страниц': 11, 'Ответы': 34, 'Ошибки': 26, 'Первый': 12, 'Последний': 12}
    data_sheet(wb, name, d.drop(columns='_bad'), 'Динамические блоки', 'Что подгружается внутри страницы: само или по действию человека', widths,
               wrap=('Блок', 'Вид', 'Ответы', 'Ошибки'), center=('Загрузка',), kpi=kpi, kpi_col='Запросов',
               row_rule=lambda r, _it=iter(bad): F_NOTE if next(_it) else None,   # блоки с ошибками — персиковым, как в «Файлах»
               links=[('Сводка по видам', 'Анатомия сайта'), ('Куда формы отправляют данные', 'Точки приёма данных'), ('Сочетания фильтров', 'Фасеты')])
    ws = wb[name]
    for i, k_ in enumerate(d['Вид']):   # неопознанные — серым; ошибки видны красным в колонке «Ошибки», много это или мало — решает аналитик
        fill = 'EFEFEF' if k_ == 'Неопознанные' else None
        if fill:
            for c in range(2, 2 + len(d.columns) - 1): ws.cell(5 + i, c).fill = PatternFill('solid', fgColor=fill)
    return ws


def pages_sheet(wb, T, name, mode):
    """«Разделы», «Типы страниц», «Страницы» — один лист на срезе реестра адресов (blocks.top_pages)."""
    if T is None or not len(T) or name not in wb.sheetnames: return None
    from urllib.parse import unquote
    from .blocks import TOP_CHANNELS
    chans = [lab for _, lab in TOP_CHANNELS] + ['Прочие']
    dec = lambda x: unquote(str(x), errors='replace')
    first = {'Разделы': 'Раздел', 'Типы страниц': 'Тип страницы', 'Страницы': 'Страница'}[mode]
    cols = {first: T['страница'].map(dec)}
    if mode == 'Разделы': cols['Статус'] = T['статус']
    if mode in ('Разделы', 'Типы страниц'): cols['Страниц'] = T['страниц'].astype(int)
    cols.update({'Просмотров': T['просмотров'].astype(int), 'Визитов (люди)': T['визитов'].astype(int), 'Уникальных IP (люди)': T['ip'].astype(int),
                 'Заявок отправлено': T['заявок'].astype(int)})
    if 'ответы' in T:
        sp = T['ответы'].fillna('').map(split_codes)
        cols['Ответы'], cols['Ошибки'] = sp.str[0], sp.str[1]
    else:
        cols['Ошибки'] = T['ошибки'].map(lambda x: split_codes(x)[1] if x else '')
    cols.update({c_: T[c_].astype(int) for c_ in chans})
    if mode == 'Типы страниц': cols['Пример'] = T['пример'].map(dec)
    d = pd.DataFrame(cols)
    dead = list(T['статус'].isin(['не существует', 'сломан'])) if mode == 'Разделы' else list(d['Ошибки'].astype(bool))   # на страницах — строки с ошибками
    kpi = [({'Разделы': 'Разделов', 'Типы страниц': 'Типов', 'Страницы': 'Страниц'}[mode], len(d))]
    if mode == 'Разделы': kpi.append(('Не существуют', int(sum(dead))))
    kpi += [('Визитов (люди)', int(d['Визитов (люди)'].sum())), ('Заявок', int(d['Заявок отправлено'].sum()))]
    if mode != 'Разделы': kpi.append(('С ошибками', int(d['Ошибки'].astype(bool).sum())))
    widths = {first: 55 if mode != 'Разделы' else 30, 'Статус': 14, 'Страниц': 9, 'Просмотров': 11, 'Визитов (люди)': 12, 'Уникальных IP (люди)': 11,
              'Заявок отправлено': 11, 'Ответы': 30, 'Ошибки': 22, 'Пример': 50, **{c_: 11 for c_ in chans}}
    note = {'Разделы': 'Запрашиваемые разделы сайта', 'Типы страниц': 'Страницы одного вида, собранные по шаблону адреса',
            'Страницы': 'TOP500 — самые посещаемые страницы сайта'}[mode]
    links = {'Разделы': [('Типы страниц внутри разделов', 'Типы страниц'), ('Устройство разделов', 'Анатомия сайта')],
             'Типы страниц': [('Разделы', 'Разделы'), ('Отдельные страницы', 'Страницы'), ('Каталоги по уровням', 'Анатомия сайта')],
             'Страницы': [('Типы страниц', 'Типы страниц'), ('Разделы', 'Разделы')]}[mode]
    data_sheet(wb, name, d, mode, note, widths, wrap=(first, 'Ответы', 'Ошибки', 'Пример'), center=('Статус',),
               kpi=kpi, kpi_col='Просмотров', row_rule=lambda r, _it=iter(dead): F_NOTE if next(_it) else None, links=links)
    ws = wb[name]
    r_ = ws.max_row + 2
    notes = {'Разделы': ['Раздел — первый уровень адреса страниц. «Не существует» — людям там ни разу не ответили 2xx.'],
             'Типы страниц': ['Тип страницы — адреса, которые различаются только номерами и названиями (* — изменяемая часть). Пример — самая посещаемая страница этого типа.'],
             'Страницы': ['Только страницы, которые отвечали людям 200.']}[mode]
    notes.append('Заявок отправлено — отправки форм с этих страниц. Каналы — с чего начался визит. Зонды, файлы и битые адреса — на своих листах.')
    for t_ in notes:
        c_ = ws.cell(r_, 2, t_); c_.font = Font(name='Arial', size=9, italic=True, color=GREY); r_ += 1
    return ws


WEEKDAY = ('Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс')


def journal_sheet(wb, name, A, subtitle, kpi, links, notes, top_labels=None, kpi_col=None, outages=None):
    """Журнал по дням с двухуровневой шапкой: «День» (день недели, дата, время в логе) и группы колонок вида «Верх|Низ».
    A — строки по дням (день, _с, _по, _полный, «Верх|Низ»…) и строка «Итого». Общий для «Журнала активности» и «Журнала ошибок»."""
    if A is None or not len(A) or name not in wb.sheetnames: return None
    top_labels = top_labels or {}
    body = A[A['день'] != 'Итого']
    tot = A[A['день'] == 'Итого']
    dt_ = pd.to_datetime(body['день'])
    groups = [c_ for c_ in A.columns if '|' in c_]
    cols = {'День': [WEEKDAY[d.weekday()] for d in dt_], 'Дата': dt_.dt.strftime('%d.%m.%Y').tolist(),
            'Время': ['00:00–23:59' if full else f'{a}–{b}' for a, b, full in zip(body['_с'], body['_по'], body['_полный'])]}
    for g in groups: cols[g] = body[g].astype(int).tolist()
    d = pd.DataFrame(cols)
    if len(tot):
        t_ = {'День': 'Итого', 'Дата': '', 'Время': f'{len(body)} дн.'}
        t_.update({g: int(tot[g].iloc[0]) for g in groups})
        d = pd.concat([d, pd.DataFrame([t_])], ignore_index=True)
    full_days = body['_полный'].values
    weekend = [x >= 5 for x in dt_.dt.weekday] + [False] * len(tot)
    od = {}
    if outages is not None and len(outages):   # дни со сбоями: строка персиковая, сбои — последней колонкой (порядок остальных колонок не меняется)
        for _, o in outages.iterrows():
            od.setdefault(o['день'], []).append(f"{o['сбой']}, {int(o['минут_всего'])} мин")
        d['Сбой'] = ['\n'.join(od.get(x, [])) for x in body['день']] + [''] * len(tot)
    row_fill = [F_NOTE if (i < len(body) and body['день'].iloc[i] in od) else ('F3F3F3' if w else None) for i, w in enumerate(weekend)]
    widths = {'День': 7, 'Дата': 12, 'Время': 13, 'Сбой': 20, **{g: 10 for g in groups}}
    data_sheet(wb, name, d, name, subtitle, widths, center=('День', 'Дата', 'Время'), wrap=('Сбой',),
               kpi=kpi(body) if callable(kpi) else kpi, kpi_col=kpi_col if kpi_col in groups else None,
               row_rule=lambda r, _it=iter(row_fill): next(_it), links=links)
    ws = wb[name]
    H, X = 4, 1
    # шапка в два уровня: строка 3 — группы (объединённые ячейки), строка 4 — колонки
    hdr = ['День', 'День', 'День'] + [top_labels.get(g.split('|')[0], g.split('|')[0]) for g in groups] + (['Сбой'] if od else [])
    j = 0
    while j < len(hdr):
        k = j
        while k + 1 < len(hdr) and hdr[k + 1] == hdr[j]: k += 1
        c_ = ws.cell(H - 1, j + 1 + X, hdr[j])
        c_.font = Font(name='Arial', size=9, bold=True, color='FFFFFF'); c_.fill = PatternFill('solid', fgColor=ORANGE)
        c_.alignment = Alignment(horizontal='center', vertical='center')
        for q in range(j, k + 1):
            cc = ws.cell(H - 1, q + 1 + X); cc.fill = PatternFill('solid', fgColor=ORANGE)
            cc.border = Border(left=WHITE if q == j else None, right=WHITE if q == k else None, bottom=WHITE)
        if k > j: ws.merge_cells(start_row=H - 1, start_column=j + 1 + X, end_row=H - 1, end_column=k + 1 + X)
        j = k + 1
    ws.row_dimensions[H - 1].height = 20
    for q, g in enumerate(groups):
        ws.cell(H, q + 4 + X).value = g.split('|')[1]
    if od:   # «Сбой» — одна колонка на два уровня шапки
        col_ = len(hdr) + X
        ws.cell(H, col_).value = None
        ws.merge_cells(start_row=H - 1, start_column=col_, end_row=H, end_column=col_)
    ws.row_dimensions[H].height = 20
    # неполные дни — серым курсивом; «Итого» — жирным с линией сверху
    for i, full in enumerate(full_days):
        if not full:
            for q in range(len(d.columns)):
                c_ = ws.cell(H + 1 + i, q + 1 + X); c_.font = Font(name='Arial', size=9, italic=True, color=GREY)
    if len(tot):
        r_ = H + len(d)
        for q in range(len(d.columns)):
            c_ = ws.cell(r_, q + 1 + X); c_.font = Font(name='Arial', size=9, bold=True)
            c_.border = Border(top=Side(style='thin', color=ORANGE), bottom=SEP)
        ws.auto_filter.ref = f"{ws.cell(H, 1 + X).coordinate}:{ws.cell(H + len(d) - 1, len(d.columns) + X).coordinate}"
    r_ = ws.max_row + 2
    for t_ in notes:
        c_ = ws.cell(r_, 2, t_); c_.font = Font(name='Arial', size=9, italic=True, color=GREY); r_ += 1
    legend(ws, (['сбой'] if od else []) + ['выходные', 'неполный'])
    return ws


def _day(ts): return f"{WEEKDAY[pd.Timestamp(ts).weekday()]} {pd.Timestamp(ts).strftime('%d.%m')}"


def activity_sheet(wb, A, name='Журнал активности', outages=None):
    """Журнал активности: визиты, уникальные IP и заявки по дням для людей, роботов, ботов и своих."""
    def kpi(body):
        hv = body.loc[body['_полный'], 'Визиты|Люди'] if body['_полный'].any() else body['Визиты|Люди']
        return [('Визиты людей в день', int(round(hv.mean())) if len(hv) else 0),
                ('Самый активный день', _day(body.loc[body['Визиты|Люди'].idxmax(), 'день']) if len(body) else ''),
                ('Заявки людей', int(body['Заявки|Люди'].sum()))]
    return journal_sheet(wb, name, A, 'Визиты, уникальные IP и заявки по дням', kpi,
                         [('Кто заходит на сайт', 'Активность'), ('Все отправки форм', 'Конверсии')] + ([('Сбои', ('NXLD_02_Errors.xlsx', 'Сбои'))] if outages is not None and len(outages) else []),
                         ['Визиты — по группам, как на листе «Активность». IP — разные адреса за день; в «Итого» — разные за весь период, поэтому меньше суммы по дням.',
                          'Заявки — отправки форм; у ботов — спам, у своих — тесты. Время — какая часть суток попала в лог; неполные дни в среднее не входят.'],
                         top_labels={'IP': 'Уникальные IP'}, kpi_col='IP|Люди', outages=outages)


SRC_LABEL = {'документация': 'Документация', 'сообщество': 'Сообщество', 'сборник': 'Сообщество', 'наблюдение': 'Сообщество',
             'поиск': 'Оперативный поиск', 'поиск, подтверждено': 'Оперативный поиск, подтверждено', 'дедукция': 'Дедукция', 'поведение': 'Дедукция', 'имя': 'Дедукция', '': ''}


def params_sheet(wb, res, name='Параметры запросов'):
    """Каждый ключ (семейство ключей) после «?»: группа, что это, откуда знаем, запросы, люди, значения, где встречается."""
    P = res.get('params')
    if not P: return None
    from .anatomy import PARAM_ORDER as order
    d = pd.DataFrame([{'Параметр': p['параметр'], 'Группа': p['группа'], 'Запросов': int(p['запросов']), 'Визитов (люди)': int(p['людей']),
                       'Опознание': p.get('что', ''), 'Основание': SRC_LABEL.get(p.get('источник', ''), p.get('источник', '')),
                       'Значений': int(p['значений']), 'Частое значение': p['частое_значение'],
                       'Точки обращения': _who(p['где']), 'Спутники': p.get('вместе_с', '')} for p in P])
    if not d['Спутники'].astype(bool).any(): d = d.drop(columns='Спутники')
    cnt = d['Группа'].value_counts()
    kpi = [('Параметров', len(d))] + [(g.split(',')[0], int(cnt[g])) for g in order if cnt.get(g)]
    kpi = [k for k in kpi if k[1]] + [k for k in kpi if not k[1]]
    if name not in wb.sheetnames: wb.create_sheet(name)
    widths = {'Параметр': 30, 'Группа': 24, 'Запросов': 11, 'Визитов (люди)': 11, 'Опознание': 50, 'Основание': 18, 'Значений': 10, 'Частое значение': 36, 'Точки обращения': 40, 'Спутники': 40}
    data_sheet(wb, name, d, 'Параметры запросов', 'Что передают в адресе после «?»', widths, wrap=('Параметр', 'Группа', 'Опознание', 'Основание', 'Частое значение', 'Точки обращения', 'Спутники'),
               kpi=kpi, kpi_col='Запросов', row_rule=lambda r: F_NOTE if r.get('Группа') == 'Атаки и зонды' else 'EFEFEF' if r.get('Группа') == 'Неизвестные' else 'F7F7F7' if r.get('Основание') == 'Дедукция' else None,
               links=[('Сводка по группам', 'Анатомия сайта'), ('Фильтры каталога по значениям', 'Фасеты')])
    ws = wb[name]
    return ws


ERR_TOKEN = None


_PCT = None


def human_url(v):
    """Раскодирует в адресе только не-латинские буквы (кириллицу и т. п.) и пробел. Закодированная латиница
    (%27, %3C, %65) остаётся как в логе — это улика: так прячут атаки и зонды."""
    import re
    global _PCT
    if not isinstance(v, str) or '%' not in v: return v
    if _PCT is None: _PCT = re.compile(r'(?:%[0-9A-Fa-f]{2})+')
    def rep(m):
        raw = m.group(0)
        toks = [raw[i:i + 3] for i in range(0, len(raw), 3)]
        bs = bytes(int(t[1:], 16) for t in toks)
        out, i = [], 0
        while i < len(bs):
            b = bs[i]
            if b < 0x80:
                out.append(' ' if b == 0x20 else toks[i]); i += 1; continue
            n = 2 if b >> 5 == 0b110 else 3 if b >> 4 == 0b1110 else 4 if b >> 3 == 0b11110 else 0
            try:
                if not n: raise ValueError
                out.append(bs[i:i + n].decode('utf-8')); i += n
            except Exception:
                out.append(toks[i]); i += 1
        return ''.join(out)
    return _PCT.sub(rep, v)


def humanize_urls(wb, skip=('_snapshot',)):
    """Во всех листах: кириллица в адресах — буквами, а не %D0%BA…; по правилу human_url."""
    for ws in wb.worksheets:
        if ws.title in skip: continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and '%' in v:
                    h = human_url(v)
                    if h != v: c.value = h


def redden_codes(wb, skip=('_snapshot',), only=None, rx=r'код|ответ'):
    """Во всех листах: в колонках с кодами ответа ошибки (4xx/5xx, кроме 499) — красным #C00000, остальное как было."""
    import re
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    tok = re.compile(r'(?<!\d)([1-5]\d\d)(\s*(?::\s*[\d\s\u00a0]+|\([\d\s\u00a0]+\)))')
    for ws in wb.worksheets:
        if ws.title in skip or ws.max_row < 2 or (only is not None and ws.title not in only): continue
        hdr_rows = [r for r in range(1, min(ws.max_row, 6) + 1)]
        cols = {}
        for r in hdr_rows:
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and re.search(rx, v, re.I) and v.strip() not in ('Ответы',):
                    cols[c] = r
        for c, hr in cols.items():
            for r in range(hr + 1, ws.max_row + 1):
                cell = ws.cell(r, c); v = cell.value
                if not isinstance(v, str) or not tok.search(v): continue
                parts, pos, any_bad = [], 0, False
                f0 = cell.font
                base = InlineFont(rFont=f0.name or 'Arial', sz=f0.sz or 9, b=f0.b, color=(f0.color.rgb if f0.color is not None and isinstance(f0.color.rgb, str) else None))
                redf = InlineFont(rFont=f0.name or 'Arial', sz=f0.sz or 9, b=f0.b, color=RED)
                for m in tok.finditer(v):
                    k = int(m.group(1)); bad = k >= 400 and k != 499
                    if not bad: continue
                    any_bad = True
                    if m.start() > pos: parts.append(TextBlock(base, v[pos:m.start()]))
                    parts.append(TextBlock(redf, m.group(0))); pos = m.end()
                if any_bad:
                    if pos < len(v): parts.append(TextBlock(base, v[pos:]))
                    cell.value = CellRichText(parts)


def heat_sheet(wb, name, title, subtitle, blocks, intro=(), links=(), unit='Запросов за час'):
    """Тепловые карты «день × час» друг под другом на одном листе: blocks — [(заголовок, подпись, DataFrame день × 0..23)].
    Перед картами — заметка, куда смотреть (intro). Общий для «Нагрузки по часам» (03) и «Признаков торможения» (02)."""
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter as L_
    idx = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name, idx)
    ws.column_dimensions['A'].width = 2.5
    ws['B1'] = title.upper(); ws['B1'].font = Font(name='Montserrat', size=16, bold=True, color=ORANGE); ws.row_dimensions[1].height = 34
    ws['B2'] = subtitle; ws['B2'].font = Font(name='Comfortaa', size=11, bold=True, color=GREY); ws.row_dimensions[2].height = 24
    r = 4
    for t_ in intro:   # заметка перед картами — персиковая плашка на всю ширину
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=28)
        c_ = ws.cell(r, 2, t_); c_.font = Font(name='Arial', size=10, italic=True, color=INK); c_.fill = PatternFill('solid', fgColor=F_NOTE)
        c_.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
        ws.row_dimensions[r].height = 15 * (len(t_) // 170 + 1) + 10
        r += 1
    r += 1
    for bt, bn, P in blocks:
        if P is None or not len(P): continue
        ws.cell(r, 2, bt.upper()).font = Font(name='Montserrat', size=12, bold=True, color=ORANGE); ws.row_dimensions[r].height = 22; r += 1
        if bn:
            ws.cell(r, 2, bn).font = Font(name='Arial', size=9, italic=True, color=GREY); r += 1
        hdr = ['День', 'Дата'] + [f'{h:02d}' for h in range(24)] + ['Всего']
        for j, v in enumerate(hdr):
            c_ = ws.cell(r, 2 + j, v); c_.font = Font(name='Arial', size=9, bold=True, color='FFFFFF'); c_.fill = PatternFill('solid', fgColor=ORANGE)
            c_.alignment = Alignment(horizontal='center', vertical='center'); c_.border = Border(left=WHITE, right=WHITE)
        r += 1; r0 = r
        for day, row in P.iterrows():
            d_ = pd.Timestamp(day)
            ws.cell(r, 2, WEEKDAY[d_.weekday()]).alignment = Alignment(horizontal='center')
            ws.cell(r, 3, d_.strftime('%d.%m')).alignment = Alignment(horizontal='center')
            for h in range(24):
                c_ = ws.cell(r, 4 + h, int(row.get(h, 0))); c_.number_format = NUM_FMT; c_.font = Font(name='Arial', size=8)
            c_ = ws.cell(r, 28, int(row.sum())); c_.number_format = NUM_FMT; c_.font = Font(name='Arial', size=9, bold=True)
            for q in (2, 3): ws.cell(r, q).font = Font(name='Arial', size=9, color='000000' if d_.weekday() < 5 else GREY)
            r += 1
        ws.conditional_formatting.add(f'D{r0}:AA{r - 1}', ColorScaleRule(start_type='min', start_color='FFFFFF', mid_type='percentile', mid_value=50,
                                                                        mid_color='FCD9C4', end_type='max', end_color='E8541C'))
        r += 1
    ws.column_dimensions['B'].width = 6; ws.column_dimensions['C'].width = 7
    for j in range(24): ws.column_dimensions[L_(4 + j)].width = 6.3
    ws.column_dimensions['AB'].width = 10
    for text, target in links:
        c = ws.cell(r, 2, f'{text}: лист «{target}» →'); c.hyperlink = f"#'{target}'!A1"; c.font = Font(name='Arial', size=10, color=ORANGE2, underline='single'); r += 1
    c_ = ws.cell(r + 1, 2, f'{unit}. Цвет — от белого (меньше всего) до тёмно-оранжевого (больше всего) в каждой карте отдельно. Выходные — серой датой.')
    c_.font = Font(name='Arial', size=9, italic=True, color=GREY)
    ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    return ws


def extra_table(ws, df, title, note='', widths=None, wrap=(), row_rule=None, size=9, at=None):
    """Вторая таблица на листе — под первой (после ссылок и заметок): заголовок, подпись, оранжевая шапка, строки.
    at — буквы колонок листа для каждой колонки таблицы (широкие колонки — на широкие места первой таблицы)."""
    from openpyxl.utils import column_index_from_string as ci_
    r = ws.max_row + 3
    ws.cell(r, 2, title.upper()).font = Font(name='Montserrat', size=12, bold=True, color=ORANGE); ws.row_dimensions[r].height = 22; r += 1
    if note:
        ws.cell(r, 2, note).font = Font(name='Arial', size=9, italic=True, color=GREY); r += 1
    cols = list(df.columns)
    pos = [ci_(x) for x in at] if at else [2 + j for j in range(len(cols))]
    for j, c in enumerate(cols):
        cell = ws.cell(r, pos[j], c); cell.font = Font(name='Arial', size=size, bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=ORANGE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = Border(left=WHITE, right=WHITE)
    ws.row_dimensions[r].height = 32
    for row in df.itertuples(index=False):
        r += 1
        rf = row_rule(dict(zip(cols, row))) if row_rule else None
        for j, v in enumerate(row):
            if isinstance(v, float) and pd.isna(v): v = None
            if hasattr(v, 'item'): v = v.item()
            if isinstance(v, str): v = ILLEGAL.sub('�', v)
            cell = ws.cell(r, pos[j], v)
            isnum = isinstance(v, numbers.Number) and not isinstance(v, bool)
            cell.font = Font(name='Arial', size=size)
            cell.alignment = Alignment(horizontal='right' if isnum else 'left', vertical='top', wrap_text=cols[j] in wrap, indent=1)
            if isnum and isinstance(v, int) and abs(v) >= 1000: cell.number_format = NUM_FMT
            cell.border = Border(bottom=SEP)
            if rf: cell.fill = PatternFill('solid', fgColor=rf)
    return ws
