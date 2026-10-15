"""#22: для магазина основная конверсия — заказ: разные номера на странице успеха оформления. Пересчёты оформления, проверка подписки, кабинет и админка — не цели."""
import os
import tempfile
from datetime import timedelta

import pandas as pd

import synth
from nxld import analyze, prepare

AJAX = '/bitrix/components/bitrix/sale.order.ajax/ajax.php'
SUBSCRIBE = '/bitrix/components/bitrix/catalog.product.subscribe/ajax.php'
OK = 'https://site.ru/personal/order/make/'


def _log(shop=True, calls=(3, 7, 10, 4, 5)):
    log = synth.Log()
    for k in range(30):   # люди: движок определяется по 20 и больше IP
        ip = f'95.165.{k // 200}.{k + 1}'
        log.visit(ip, synth.T0 + timedelta(hours=8, minutes=9 * k), ref='https://yandex.ru/')
    for k in range(4):   # форма обратной связи — цель, как раньше
        ip = f'95.166.0.{k + 1}'
        t = synth.T0 + timedelta(days=1, hours=9, minutes=5 * k)
        log.visit(ip, t)
        log.line(ip, t + timedelta(seconds=100), '/contacts/send/', 200, ref='https://site.ru/contacts/', method='POST')
    if not shop: return log
    for n, calls_ in enumerate(calls, 1):   # 5 заказов: на каждый 3–10 пересчётов оформления, проверки подписки, кабинет
        ip = f'95.167.0.{n}'
        t = synth.T0 + timedelta(days=2, hours=10, minutes=20 * n)
        log.visit(ip, t)
        for j in range(calls_): log.line(ip, t + timedelta(seconds=60 + 3 * j), AJAX, 200, ref='https://site.ru/personal/order/make/', method='POST')
        log.line(ip, t + timedelta(seconds=40), SUBSCRIBE, 200, ref='https://site.ru/catalog/', method='POST')
        log.line(ip, t + timedelta(seconds=400), f'/personal/order/make/?ORDER_ID={1000 + n}', 200, ref='https://site.ru/personal/order/make/')
        log.line(ip, t + timedelta(seconds=460), f'/personal/order/make/?ORDER_ID={1000 + n}', 200, ref='https://site.ru/personal/order/make/')   # обновил страницу успеха
        log.line(ip, t + timedelta(seconds=520), '/personal/order/', 200, ref='https://site.ru/', method='POST')   # кабинет
    return log


def _prepare(log):
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        log.write(path)
        prepare.run([path], work, log=lambda *a: None)
        return pd.read_pickle(os.path.join(work, 'R.pkl')), pd.read_pickle(os.path.join(work, 'V.pkl'))


def test_orders_counted_by_number():
    R, V = _prepare(_log())
    g = R['goal'].astype(str)
    assert int(R['is_order'].sum()) == 5   # пять номеров, обновление страницы успеха не считается
    assert int((g == '/contacts/send/').sum()) == 4 and int(g.str.contains('sale.order.ajax|subscribe|/personal/order/$|admin', regex=True).sum()) == 0   # форма — цель; шаги, подписка, кабинет — нет
    assert int(R['goal_success'].sum()) == 9   # 5 заказов + 4 заявки
    assert int(V['n_conv'].sum()) == 9 and int(V['n_goal'].sum()) == 9
    assert int(R['order_step'].sum()) == 3 + 7 + 10 + 4 + 5   # шаги оформления отмечены для карточки о потерянных заявках


def test_order_counted_once_in_first_visit():
    log = _log(calls=(3,) * 5)
    ip = '95.167.0.1'
    log.line(ip, synth.T0 + timedelta(days=5, hours=10), '/personal/order/make/?ORDER_ID=1001', 200, ref='-')   # позже тот же номер из другого визита
    R, V = _prepare(log)
    assert int(R['is_order'].sum()) == 5
    first = R.loc[R['is_order'], 'vid']
    assert first.is_unique   # по заказу — свой визит, номер в позднем визите не посчитан


def test_shop_without_orders_and_forms_as_before():
    """Сайт без магазина: нет страниц успеха заказа — цели только по формам, как раньше."""
    R, V = _prepare(_log(shop=False))
    assert int(R['is_order'].sum()) == 0 and int(R['goal_success'].sum()) == 4 and int(V['n_conv'].sum()) == 4


def test_analyze_accepted_count_follows_orders(monkeypatch):
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        _log().write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
        s = res['summary']['Общий анализ']
        assert s['Принято от людей'] == 9 and s['Отправок целей всего'] == 9, s


def _lost_log(win_ok=True):
    """3 ошибки 504 на оформлении, после двух — заказ в течение часа; третий покупатель заказа не оформил."""
    log = _log(calls=(3, 3, 3, 3, 3))
    for k in range(3):
        ip = f'95.168.0.{k + 1}'
        t = synth.T0 + timedelta(days=3, hours=10, minutes=30 * k)
        log.visit(ip, t)
        log.line(ip, t + timedelta(seconds=60), AJAX, 504, ref='https://site.ru/personal/order/make/', method='POST')
        if k < 2: log.line(ip, t + timedelta(minutes=40), f'/personal/order/make/?ORDER_ID={2000 + k}', 200, ref='https://site.ru/personal/order/make/')
    return log


def _card(log, monkeypatch):
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        log.write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
        return [x for x in res['findings'] if x['key'].split(':')[1] == 'lost_leads']


def test_lost_leads_two_numbers_main_is_without_order(monkeypatch):
    card = _card(_lost_log(), monkeypatch)
    assert len(card) == 1, card
    c0 = card[0]
    assert c0['главная_цифра'] == 1   # потеряно 1 из 3
    assert 'Ошибок всего 3, из них закончились заказом 2, без заказа 1' in c0['факты'], c0['факты']


def test_lost_leads_all_ended_with_orders_no_card(monkeypatch):
    log = _log(calls=(3,) * 5)
    for k in range(2):
        ip = f'95.168.0.{k + 1}'
        t = synth.T0 + timedelta(days=3, hours=10, minutes=30 * k)
        log.visit(ip, t)
        log.line(ip, t + timedelta(seconds=60), AJAX, 504, ref='https://site.ru/personal/order/make/', method='POST')
        log.line(ip, t + timedelta(minutes=50), f'/personal/order/make/?ORDER_ID={3000 + k}', 200, ref='https://site.ru/personal/order/make/')
    assert _card(log, monkeypatch) == []


def test_lost_leads_forms_without_number_as_before(monkeypatch):
    log = _log(shop=False)
    for k in range(2):   # форма обратной связи отвечает ошибкой, номера нет — потеря, как раньше
        ip = f'95.169.0.{k + 1}'
        t = synth.T0 + timedelta(days=3, hours=10, minutes=30 * k)
        log.visit(ip, t)
        log.line(ip, t + timedelta(seconds=60), '/contacts/send/', 500, ref='https://site.ru/contacts/', method='POST')
    card = _card(log, monkeypatch)
    assert len(card) == 1 and card[0]['главная_цифра'] == 2 and 'Ошибок всего' not in card[0]['факты']


def test_order_with_error_on_success_page_counted_and_flagged(monkeypatch):
    """B1: номер заказа на странице успеха с ответом 403 — заказ засчитан, есть карточка с кодом."""
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    log = _log()
    ip = '95.167.0.9'
    t = synth.T0 + timedelta(days=2, hours=15)
    log.visit(ip, t)
    log.line(ip, t + timedelta(seconds=400), '/personal/order/make/?ORDER_ID=1999', 403, ref='https://site.ru/personal/order/make/')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        log.write(path)
        prepare.run([path], work, log=lambda *a: None)
        assert int(pd.read_pickle(os.path.join(work, 'R.pkl'))['is_order'].sum()) == 6
        res = analyze.run(work, log=lambda *a: None)
    card = [x for x in res['findings'] if x['key'].split(':')[1] == 'order_page_error']
    assert len(card) == 1 and '403 (1)' in card[0]['факты'] and 'из 6' in card[0]['факты'], card


def test_no_error_card_when_success_page_ok(monkeypatch):
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        _log().write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
    assert not [x for x in res['findings'] if x['key'].split(':')[1] == 'order_page_error']


def _funnel_log():
    log = _log()
    for k in range(25):   # 25 людей добавляют товар: POST на один обработчик со страниц каталога
        ip = f'95.165.0.{k + 1}'
        log.line(ip, synth.T0 + timedelta(hours=8, minutes=9 * k, seconds=30), '/catalog/add/', 200, ref=f'https://site.ru/catalog/p{k}/', method='POST')
    for n in range(1, 6):   # все 5 заказов прошли корзину (с 3 ниже — 8 корзин) и оформление
        ip, t = f'95.167.0.{n}', synth.T0 + timedelta(days=2, hours=10, minutes=20 * n)
        log.line(ip, t + timedelta(seconds=20), '/personal/cart/', 200, ref='https://site.ru/catalog/')
        log.line(ip, t + timedelta(seconds=30), '/personal/order/make/', 200, ref='https://site.ru/personal/cart/')
    for k in range(3):   # корзина и оформление без заказа
        ip, t = f'95.168.0.{k + 1}', synth.T0 + timedelta(days=3, hours=10, minutes=20 * k)
        log.visit(ip, t)
        log.line(ip, t + timedelta(seconds=20), '/personal/cart/', 200, ref='https://site.ru/catalog/')
        log.line(ip, t + timedelta(seconds=30), '/personal/order/make/', 200, ref='https://site.ru/personal/cart/')
    return log


def test_order_funnel_steps(monkeypatch):
    """B2: 5 заказов, 8 корзин, 3 визита с оформлением без заказа → числа шагов воронки."""
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        _funnel_log().write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
    of = res['sheets']['Маркетинг']['Воронка заказа']
    st = dict(of['шаги'])
    assert st['order'] == 5 and st['cart'] == 8 and st['checkout'] == 8 and st['action'] == 25, st
    assert st['people'] >= 41 and of['заказ_через_корзину'] == 5 and of['действия_определены']


def test_order_funnel_absent_without_shop(monkeypatch):
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work = os.path.join(d, 'access.log'), os.path.join(d, 'work')
        _log(shop=False).write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
    assert res['sheets']['Маркетинг']['Воронка заказа'] is None


def test_order_funnel_on_funnels_sheet(tmp_path):
    """B2: на листе «Воронки» файла 05 — воронка заказа с подписями из data/locale/ru/funnel.json и примечанием; строк «Заказ» среди форм нет."""
    import glob
    import openpyxl
    _, out, _ = synth.run(_funnel_log(), str(tmp_path))
    f = glob.glob(os.path.join(out, 'NXLD_05*.xlsx'))[0]
    ws = openpyxl.load_workbook(f)['Воронки']
    cells = [(c.row, c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    texts_ = [str(v) for _, v in cells]
    assert 'ВОРОНКА ЗАКАЗА' in texts_
    for lab in ('Визиты людей', 'Действия с товаром', 'Корзина', 'Оформление', 'Заказ'): assert lab in texts_, lab
    row = {v: r for r, v in cells if v in ('Корзина', 'Оформление', 'Заказ')}
    vals = {k: [c.value for c in ws[row[k]] if c.value is not None] for k in row}
    assert vals['Корзина'][1:3] == [8, vals['Корзина'][2]] and vals['Оформление'][1] == 8 and vals['Заказ'][1] == 5 and vals['Заказ'][2] == round(5 / 8 * 100, 1), vals
    assert any('Из визитов с заказом (5) были в корзине' in t and t.rstrip('.').endswith('5)') for t in texts_)
    assert not any('/personal/order/make/?ORDER_ID' in t for t in texts_)   # строка цели «Заказ» среди форм заменена воронкой
