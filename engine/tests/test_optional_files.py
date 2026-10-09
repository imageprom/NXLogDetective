"""#10: 404 на необязательный стандартный файл — норма: ни карточки, ни рекомендации, серая строка на «Файлах для роботов»;
файлы для ИИ — мягкое «Замечание»; тот же файл с ответом 500 — карточка."""
import glob
import json
import os
from datetime import timedelta

import synth

OPT = ['/.well-known/traffic-advice', '/.well-known/passkey-endpoints', '/.well-known/change-password', '/.well-known/security.txt',
       '/security.txt', '/ads.txt', '/app-ads.txt', '/humans.txt', '/llms.txt', '/llms-full.txt', '/ai.txt']
UAS = [synth.UA_CHROME, 'Mozilla/5.0 (compatible; YandexBot/3.0; +http://yandex.com/bots)', 'Mozilla/5.0 (compatible; GPTBot/1.1; +https://openai.com/gptbot)']


def _log(code=404, only=None):
    log = synth.Log()
    log.people()
    for d in range(7):
        for k, u in enumerate(OPT):
            for j, ua in enumerate(UAS):
                c = code if only is None or u == only else 404
                log.line(f'203.0.113.{j * 20 + k + 1}', synth.T0 + timedelta(days=d, hours=6 + j, minutes=k), u, c, ua=ua)
    return log


def test_optional_404_is_norm(tmp_path):
    res, out, _ = synth.run(_log(), str(tmp_path))
    kinds = {x['key'].split(':')[1]: x for x in res['findings']}
    assert not any(k in kinds for k in ('no_service', 'service_err', 'missing_static')), sorted(kinds)
    assert not [x for x in res['findings'] if any(f in x['key'] for f in OPT)], [x['key'] for x in res['findings']]
    ai = kinds['ai_index']
    assert ai['важность'] == 'Замечание' and 'практической пользы для большинства сайтов нет' in (ai['что_сделать'] + ai.get('факт', '') + ai.get('факты', '')), ai
    import openpyxl
    for f in glob.glob(os.path.join(out, '*.xlsx')):
        for ws in openpyxl.load_workbook(f, read_only=True).worksheets:
            for row in ws.iter_rows(values_only=True):
                assert not any(isinstance(v, str) and 'Можно добавить' in v for v in row), (os.path.basename(f), ws.title)
    ws = openpyxl.load_workbook(glob.glob(os.path.join(out, '*_06_*.xlsx'))[0] if glob.glob(os.path.join(out, '*_06_*.xlsx'))
                                else glob.glob(os.path.join(out, '*_01_*.xlsx'))[0])['Файлы для роботов']
    grey = [r for r in ws.iter_rows(min_row=5) if any(c.value == 'необязательный, отсутствует — норма' for c in r)]
    assert len(grey) >= 5, [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
    hdr = [c.value for c in ws[4]]
    i = hdr.index('Вывод')
    fonts = {r[i].font.color.rgb if r[i].font.color is not None else None for r in grey}
    assert fonts and None not in fonts and not any(str(f).endswith('000000') for f in fonts), fonts
    O = json.load(open(os.path.join(synth.ENGINE, '..', 'data', 'reference', 'extensions.json'), encoding='utf-8'))['optional_files']
    assert O['to_verify'] and any('traffic-advice' in p for p in O['paths'])


def test_optional_500_is_card(tmp_path):
    res, _, _ = synth.run(_log(500, '/ads.txt'), str(tmp_path))
    assert any(x['key'] == 'Ошибки:service_err:/ads.txt' for x in res['findings']), [x['key'] for x in res['findings']]
