"""Мягкие ошибки (02): адреса без косой черты отдают другую страницу одного размера — одна группа, «что это», карточка."""
import os
import random
from datetime import timedelta

import synth

N = 600   # карточка — от 500 адресов в группе (в листе — от 200)


def test_soft_errors_no_slash(tmp_path):
    log = synth.Log()
    log.people()
    for i in range(N):
        ip = synth.RU[i % len(synth.RU)]
        t = synth.T0 + timedelta(days=i % 7, hours=14, minutes=(i // 7) % 600 // 10, seconds=i % 60)
        a = f'/catalog/x/item-{i}'
        log.line(ip, t, a, 200, ref='https://yandex.ru/', size=17000 + random.randint(0, 80))           # без «/» — одна и та же короткая страница
        log.line(ip, t + timedelta(seconds=20), a + '/', 200, ref=f'https://site.ru{a}', size=23000 + random.randint(0, 300))   # со «/» — настоящая
    res, out, _ = synth.run(log, str(tmp_path))
    G = res['sheets']['Ошибки']['Идентичные ответы']
    g = G[G['примеры'].str.contains('/catalog/x/item-')]
    assert len(g) == 1, G[['размер', 'разных_адресов']]
    r = g.iloc[0]
    assert r['разных_адресов'] >= N and '–' in r['размер']
    assert r['что_это'].startswith('Адреса без косой черты в конце')
    assert not G['размер'].str.startswith('23').any()   # страницы со «/» — разные адреса, но не группа «мягких ошибок»
    x = [x for x in res['findings'] if x['key'].split(':')[1] == 'soft_errors']
    assert len(x) == 1 and x[0]['блок'] == 'Ошибки' and x[0]['важность'] == 'Важно'
    assert x[0]['key'] == f"Ошибки:soft_errors:{r['_от']}"
    assert 'переадресация 301 на адрес с косой чертой' in x[0]['что_сделать'].lower()
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out, 'NXLD_02_Errors.xlsx'))
    assert 'Мягкие ошибки' in wb.sheetnames and 'Пустые ответы' not in wb.sheetnames and 'Одинаковые ответы' not in wb.sheetnames
    vals = [v for row in wb['Мягкие ошибки'].iter_rows(values_only=True) for v in row if v]
    assert 'ИДЕНТИЧНЫЕ ОТВЕТЫ' in vals and 'МЯГКИЕ ОШИБКИ' in vals   # заголовки листа и блока — прописными
    ov = [v for row in wb['Обзор'].iter_rows(values_only=True) for v in row if isinstance(v, str)]
    assert any(v.startswith('Мягкие ошибки') for v in ov)
