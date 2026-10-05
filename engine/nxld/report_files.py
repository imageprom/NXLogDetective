"""NXLD: лист «Файлы» — какие логи разобраны (ТЗ: технические данные — в общем файле, отдельным листом в конце)."""
from collections import Counter
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, PatternFill, Side
from .report_index import Sheet, ORANGE, ru_num, ru_short, plural, ru_date, NUM_FMT

F_LIGHT = 'F7F7F7'


def big_table(S, headers, rows, cols=('B', 'C', 'D', 'E', 'F'), num=(1,), center=(2, 3), wrap=(0, 4)):
    """Таблица в стиле больших таблиц (как «Конверсии»): оранжевая шапка шрифтом 9, строки без заливки, тонкий разделитель, отступ в ячейках."""
    ws = S.ws
    white, sep = Side(style='thin', color='FFFFFF'), Side(style='thin', color='D9D9D9')
    for i, h in enumerate(headers):
        c = ws[f'{cols[i]}{S.r}']; c.value = h
        c.font = Font(name='Arial', size=9, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=ORANGE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(left=white, right=white)
    ws.row_dimensions[S.r].height = 32
    S.r += 1
    for row in rows:
        lines = 1
        for i, v in enumerate(row):
            c = ws[f'{cols[i]}{S.r}']; c.value = v
            c.font = Font(name='Arial', size=9, color='000000')
            hz = 'center' if i in center else ('right' if i in num else 'left')
            c.alignment = Alignment(horizontal=hz, vertical='top', wrap_text=i in wrap, indent=0 if hz == 'center' else 1)
            c.border = Border(bottom=sep)
            if i in num and isinstance(v, int) and v >= 1000: c.number_format = NUM_FMT
            if i in wrap and isinstance(v, str):
                w_ = max(8, int(ws.column_dimensions[cols[i]].width or 10) - 2)
                lines = max(lines, sum(max(1, -(-len(x) // int(w_ * 1.3))) for x in v.split('\n')))
        ws.row_dimensions[S.r].height = 12 * lines + 3
        S.r += 1
    S.r += 1


def dt(s):
    t = pd.Timestamp(str(s)); return f'{t.day:02d}.{t.month:02d} {t:%H:%M}'


def shown(name):
    if ':' in name:
        arc, inner = name.split(':', 1)
        return f"{inner}\nв архиве {arc}"
    return name


def build_files(wb, res, title='Логи'):
    inv = res['inventory']
    F = pd.DataFrame(inv.get('files', []))
    if not len(F): return None
    F['строк'] = pd.to_numeric(F['строк'], errors='coerce').fillna(0).astype(int)
    F['нераспознано'] = pd.to_numeric(F['нераспознано'], errors='coerce').fillna(0).astype(int)
    main = ((res.get('site_map') or {}).get('site_hosts') or [''])[0]
    ws = wb.create_sheet(title)
    S = Sheet(ws)
    ws.row_dimensions[1].height = 12
    S.r = 2
    ws.merge_cells('B2:F2')
    S.cell('B', 'ЛОГИ', Font(name='Montserrat', size=16, bold=True, color=ORANGE)); ws.row_dimensions[2].height = 30
    ws.merge_cells('B3:F3')
    S.cell('B', 'Какие файлы логов разобраны, за какой период и что в них не так', Font(name='Comfortaa', size=11, bold=True, color='666666'), row=3)
    ws.row_dimensions[3].height = 22
    S.r = 4
    dup = {d['source']: int(d.get('lines_duplicate', 0)) for d in inv.get('duplicates', []) if d.get('lines_duplicate')}
    # копии: тот же тип, период и число строк
    key = F.apply(lambda r: (r['тип'], r['с'], r['по'], r['строк']), axis=1)
    first = {}
    copy_of = {}
    for i, k in key.items():
        if k in first: copy_of[F.loc[i, 'файл']] = F.loc[first[k], 'файл']
        else: first[k] = i
    # итого
    S.section('Итого', 'Сколько файлов и строк, за какой период, что удалено как повтор и что не удалось разобрать.')
    for t, nm in (('access', 'Access-логи'), ('error', 'Error-логи')):
        g = F[F['тип'] == t]
        if not len(g): continue
        sites = Counter(g['сайт_по_имени'].astype(str))
        val = f"{len(g)} {plural(len(g), 'файл', 'файла', 'файлов')}, {ru_short(int(g['строк'].sum()))} строк" + (' (с повторами)' if t == 'access' and dup else '')
        if len(sites) > 1: val += '\n' + '\n'.join(f"{s}: {n} {plural(n, 'файл', 'файла', 'файлов')}" for s, n in sites.most_common())
        S.pair(nm, val)
    p0, p1 = inv['period']
    tz = str(inv.get('tz', ''))
    S.pair('Период', f"{ru_date(p0)} {str(p0)[11:16]} — {ru_date(p1)} {str(p1)[11:16]}" + (f"; часовой пояс {tz[:3]}:{tz[3:]}" if len(tz) == 5 else ''))
    nd = sum(dup.values())
    S.pair('Повторы на стыках', f"удалено {ru_num(nd)} {plural(nd, 'строка', 'строки', 'строк')}" if nd else 'нет')
    if copy_of: S.pair('Копии файлов', '\n'.join(f"{shown(a).splitlines()[0]} = {shown(b).splitlines()[0]}" for a, b in copy_of.items()))
    bad = int(F['нераспознано'].sum())
    S.pair('Нераспознанные строки', ru_num(bad) if bad else 'нет')
    gaps = inv.get('hour_gaps') or []
    S.pair('Пропуски по часам', '\n'.join(str(x) for x in gaps[:10]) if gaps else 'нет')
    # по типам
    for t, nm in (('access', 'Access-логи'), ('error', 'Error-логи')):
        g = F[F['тип'] == t].sort_values(['с', 'файл'])
        if not len(g): continue
        S.section(nm, {'Access-логи': 'Журналы запросов к сайту: каждый файл, его период и формат.', 'Error-логи': 'Журналы ошибок сервера: каждый файл, его период и формат.'}[nm])
        rows = []
        for _, r in g.iterrows():
            note = []
            if r['сайт_по_имени'] and r['сайт_по_имени'] != main: note.append(f"сайт: {r['сайт_по_имени']}")
            if r['файл'] in dup: note.append(f"повтор соседнего файла: удалено {ru_num(dup[r['файл']])} строк")
            if r['файл'] in copy_of: note.append(f"копия {shown(copy_of[r['файл']]).splitlines()[0]}")
            if int(r['нераспознано']): note.append(f"не распознано {ru_num(int(r['нераспознано']))} строк")
            rows.append([shown(r['файл']), int(r['строк']), dt(r['с']), dt(r['по']), '\n'.join(note) or r['формат']])
        big_table(S, ['Файл', 'Строк', 'С', 'По', 'Формат и заметки'], rows)
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws
