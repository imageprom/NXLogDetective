"""Тексты без «программ»: 01 Обзор и Активность, 03 Админка, Нагрузка по часам, Необычные запросы, 04 Виды ботов, Утилиты (round 3, п. 3)."""
import glob
import os
from datetime import timedelta

import synth

OLD = ['и другие программы', 'Программы, которые обращаются к сайту', 'затем программы', 'программы под видом браузера',
       'программа пыталась открыть', 'Программы, которые не притворяются браузером', 'Программы без браузера', 'От программ',
       'Программы и боты', 'ПРОГРАММЫ И БОТЫ']
NEW = {('01', 'Обзор'): 'и другие утилиты', ('01', 'Активность'): 'Утилиты, которые обращаются к сайту', ('03', 'Админка'): 'затем роботы и боты',
       ('03', 'Нагрузка по часам'): 'боты под видом браузера', ('03', 'Необычные запросы'): 'отправитель пытался открыть',
       ('04', 'Виды ботов'): 'Боты и утилиты, которые не притворяются браузером', ('04', 'Утилиты'): 'Утилиты без браузера'}


def test_no_programs_in_texts(tmp_path):
    log = synth.Log()
    log.people()
    for d in range(7):
        for k, u in enumerate(('/', '/catalog/', '/about/')):   # утилита
            log.line('203.0.113.7', synth.T0 + timedelta(days=d, hours=4, seconds=k), u, 200, ua='python-requests/2.31')
            log.line('203.0.113.8', synth.T0 + timedelta(days=d, hours=4, seconds=k), u, 200, ua='curl/8.1')
        log.line('198.51.100.9', synth.T0 + timedelta(days=d, hours=5), '/', 405, method='PROPFIND', ua='Mozilla/5.0 zgrab/0.x')   # необычный метод
        log.line('198.51.100.9', synth.T0 + timedelta(days=d, hours=5, seconds=3), '\\x16\\x03\\x01', 400, method='GET', ua='-')
        for k in range(3):   # админка: посторонний получает форму входа; утилита — тоже (таблица «Роботы и боты»)
            log.line('198.51.100.20', synth.T0 + timedelta(days=d, hours=6, seconds=k), '/bitrix/admin/', 200, size=5100)
            log.line('203.0.113.7', synth.T0 + timedelta(days=d, hours=6, seconds=k), '/bitrix/admin/', 200, size=5100, ua='python-requests/2.31')
    res, out, _ = synth.run(log, str(tmp_path))
    import openpyxl
    seen = {}
    for f in glob.glob(os.path.join(out, '*.xlsx')):
        wb = openpyxl.load_workbook(f, read_only=True)
        for ws in wb.worksheets:
            text = ' '.join(str(v) for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str))
            for o in OLD:
                assert o not in text, (os.path.basename(f), ws.title, o)
            seen[(os.path.basename(f)[5:7], ws.title)] = text
    checked = [k for k in NEW if k in seen and (k != ('04', 'Виды ботов') or 'ЯВНЫЕ БОТЫ' in seen[k])]   # таблица «Явные боты» — если такие боты есть
    assert len(checked) >= 5, sorted(seen)   # листы этой синтетики
    for k in checked:
        assert NEW[k] in seen[k], (k, NEW[k])


def test_admin_robots_and_bots_title(tmp_path):
    """03 Админка: заголовок таблицы — «Роботы и боты», не «Программы и боты» (#15)."""
    from test_probes import bitrix_log
    log = bitrix_log()
    for d in range(7):
        for k in range(3):
            log.line('203.0.113.7', synth.T0 + timedelta(days=d, hours=6, seconds=k), '/bitrix/admin/', 200, size=5100, ua='python-requests/2.31')
    _, out, _ = synth.run(log, str(tmp_path), '--site', 'example.com')
    import openpyxl
    ws = openpyxl.load_workbook(glob.glob(os.path.join(out, '*_03_*.xlsx'))[0], read_only=True)['Админка']
    text = ' '.join(str(v) for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str))
    assert 'РОБОТЫ И БОТЫ' in text and 'ПРОГРАММЫ' not in text.upper(), text[:1500]

