"""NXLD: файл 04 «Боты» — по тем же правилам, что 01–03.

Обзор → Проблемы → Статистика → Разыскиваются (карточки дел) → Сигнатуры → Состав дел → листы по видам ботов →
подробные списки в конце. Данные — profiles.build и profiles.extras (дела, сигнатуры, журнал, расписание, реклама),
листы расчёта блока «Боты» и реестр обращающихся (actors)."""
from . import sheets
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side
from .report_index import Sheet, Wide, brand_header, ORANGE, ORANGE2, GREY, INK, F_NOTE, F_CARD, row_height, cap, plural
from .report_tables import data_sheet, journal_sheet, legend, heat_sheet, extra_table, split_codes, nlist
from .report_problems import plaque, OLINE, F_LIGHT

F_GREY = 'EFEFEF'
ORDER = sheets.ORDER_04
LINK = sheets.LINK_04
KEEP = ('IP',)


def relink(findings):
    for x in findings:
        if x.get('блок') == 'Боты' and x.get('лист') in LINK: x['лист'] = LINK[x['лист']]


def sheets_for(S, res):
    return {'Обзор': pd.DataFrame()}


def _d(v):
    try: return pd.to_datetime(v).strftime('%d.%m.%Y')
    except Exception: return str(v or '')


def nf(x):
    try: return f'{int(round(float(x))):,}'.replace(',', ' ')
    except Exception: return str(x)


def notes(ws, lines):
    r_ = ws.max_row + 2
    for t_ in lines:
        c_ = ws.cell(r_, 2, t_); c_.font = Font(name='Arial', size=9, italic=True, color=GREY); r_ += 1


# ---------- карточки дел ----------
def case_card(S, x):
    ws = S.ws
    title = f"ДЕЛО {x['дело']} · {x['кличка']}"
    ws.merge_cells(f'B{S.r}:E{S.r}')
    S.cell('B', title, Font(name='Arial', size=11, bold=True, color=ORANGE), None, Alignment(vertical='center', wrap_text=True))
    S.cell('F', f"{x['вид']}", Font(name='Arial', size=10, bold=True, color=INK), None, Alignment(horizontal='right', vertical='center', wrap_text=True))
    for col in 'BCDEF': ws[f'{col}{S.r}'].border = Border(bottom=OLINE)
    ws.row_dimensions[S.r].height = max(26, row_height(title, 78) + 4)
    S.r += 1
    S.pair('Обвинения', '', height=20, fill=F_CARD)   # каждое обвинение — своей строкой: статья слева, сила и что именно — справа
    for i, o in enumerate(x['обвинения'], 1):
        S.pair(f"{i}. {o['статья']}", f"{o['сила']} — {o['что']}", level=2, fill=F_LIGHT)
    pr = x['приметы'] if isinstance(x['приметы'], list) else [p_.strip() for p_ in str(x['приметы']).split(';') if p_.strip()]
    S.pair('Приметы', '\n'.join(cap(p_) for p_ in pr))
    sg = x['сигнатура']
    S.pair('Сигнатура', f"{sg['id']}: {sg['правило']}\nПроверка на логе: {sg['проверка']}", fill=F_LIGHT)
    st = x['состав']   # состав — как обвинения: подзаголовок и по строке на часть, внутри части — по строке на значение
    S.pair('Состав', f"{st['IP']} IP в {st['подсетей']} {plural(st['подсетей'], 'подсети', 'подсетях', 'подсетях')}", fill=F_CARD)
    S.pair('Подсети', '\n'.join(f"{s_} ({n_} IP, {nf(r_)} запросов)" for s_, n_, r_ in st['подсети'][:5]) + (f"\nи ещё {st['подсетей'] - 5}" if st['подсетей'] > 5 else ''), level=2, fill=F_LIGHT)
    S.pair('Сети', '\n'.join(f"{k} ({nf(v)})" for k, v in list(st['сети'].items())[:6]) + (f"\nи ещё {len(st['сети']) - 6}" if len(st['сети']) > 6 else ''), level=2, fill=F_LIGHT)
    S.pair('Страны', ', '.join(f"{k} ({nf(v)})" for k, v in st['страны'].items()), level=2, fill=F_LIGHT)
    S.pair('Роли', '\n'.join(f"{k} ({nf(v)})" for k, v in st['роли'].items()), level=2, fill=F_LIGHT)
    if x.get('сообщники'):
        S.pair('Сообщники', '\n'.join(f"Дело {d} — {n}: {w}" for d, n, w in x['сообщники']), fill=F_LIGHT)
    S.pair('Активность', '\n'.join(cap(p_.strip()) for p_ in str(x['активность']).split(';') if p_.strip()), fill=F_LIGHT)   # период, объём, часы, пик — по строке
    if x.get('сбои') or x.get('всплески'):
        S.pair('Сбои и всплески', '\n'.join([f'Сбой {t}' for t in x.get('сбои', [])] + [f'Всплеск {t}' for t in x.get('всплески', [])[:6]]))
    S.pair('Ущерб', cap(x['ущерб']), fill=F_LIGHT)
    S.pair('Меры пресечения', '\n'.join(f'— {m}' for m in x['меры']))
    ws.merge_cells(f'B{S.r}:F{S.r}')
    c = S.cell('B', f"Все IP дела: лист «Состав дел», фильтр «Дело» = {x['дело']} →", align=Alignment(vertical='center', indent=1))
    c.hyperlink = "#'Состав дел'!A1"; c.font = Font(name='Arial', size=10, color=ORANGE2, underline='single')
    ws.row_dimensions[S.r].height = 20
    S.r += 2


def wanted(wb, Pf):
    D = Pf.get('дела') or []
    if not D: return None
    ws = wb.create_sheet('Разыскиваются')
    S = Sheet(ws)
    S.psize = 9   # карточек много — шрифт как в больших таблицах
    for col, w in zip('ABCDEFG', (2.5, 32, 18, 14, 16, 56, 2.5)): ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 12
    S.r = 2
    ws.merge_cells('B2:F2')
    S.cell('B', 'РАЗЫСКИВАЮТСЯ', Font(name='Montserrat', size=16, bold=True, color=ORANGE)); ws.row_dimensions[2].height = 30
    cnt = {k: sum(1 for x in D if x['важность'] == k) for k in ('Тревога', 'Срочно', 'Важно', 'К сведению')}
    ws.merge_cells('B3:F3')
    S.cell('B', "Атакующие профили: в чём обвиняем, как узнать и что делать · " + (f"Тревога — {cnt['Тревога']} · " if cnt['Тревога'] else '') + f"Приоритетные — {cnt['Срочно']} · Важные — {cnt['Важно']} · Остальные — {cnt['К сведению']}",
           Font(name='Comfortaa', size=11, bold=True, color=GREY), row=3)
    S.r = 4
    for sev in ('Тревога', 'Срочно', 'Важно', 'К сведению'):
        xs = [x for x in D if x['важность'] == sev]
        if not xs: continue
        plaque(S, sev, len(xs), [])
        for x in xs: case_card(S, x)
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def ai_sheet(wb, AI, AP, nm='ИИ-роботы', title=None):
    """ИИ-роботы и страницы, которые нейросети открывали по запросам людей (04, или «ИИ-видимость» в 06)."""
    if AI is None or not len(AI): return None
    if nm not in wb.sheetnames: wb.create_sheet(nm)
    d = pd.DataFrame({'Назначение': AI['назначение'], 'Робот': AI['робот'], 'Запросов': AI['запросов'].astype(int), 'Уникальных IP': AI['IP'].astype(int), 'Дней': AI['дней'].astype(int),
                      'Трафик, МБ': AI['МБ'], 'Ошибок': AI['ошибок'].astype(int), 'Подлинных, %': AI['подлинных_%'], 'Что читал': AI['что_читал'].map(nlist)})
    data_sheet(wb, nm, d, nm, title or 'Роботы нейросетей: обучение, поисковый индекс, ответы на запросы людей', {'Назначение': 24, 'Робот': 30, 'Что читал': 50}, wrap=('Назначение', 'Робот', 'Что читал'),
               kpi=[('Роботов', len(d)), ('Запросов', int(d['Запросов'].sum()))], kpi_col='Дней', red=())
    if AP is not None and len(AP):
        extra_table(wb[nm], pd.DataFrame({'Страница': AP['base'], 'Робот': AP['fam'], 'Запросов': AP['запросов'].astype(int), 'Дней': AP['дней'].astype(int)}).head(200),
                    'Страницы по запросам людей', 'Что нейросети открывали, отвечая людям', wrap=('Страница',))
    return wb[nm]


def build_sheets(wb, res, S):
    Pf = res.get('profiles') or {}
    EX = res.get('bots_extra') or {}
    A = res.get('actors') or {}
    E = res.get('errors') or {}
    # Статистика
    J = EX.get('журнал')
    if J is not None and len(J):
        if 'Статистика' not in wb.sheetnames: wb.create_sheet('Статистика')
        def kpi(body):
            full = body[body['_полный'].astype(bool)] if body['_полный'].astype(bool).any() else body
            return [('Визитов ботов в день', int(round(full['Визиты ботов|Все'].mean()))), ('Новых IP за период', int(J[J['день'] == 'Итого']['Новые IP|Ботов'].iloc[0])),
                    ('Принято заявок от ботов', int(body['Заявки|Принято'].sum()))]
        journal_sheet(wb, 'Статистика', J, 'Боты по дням: виды, новые IP и заявки', kpi, [('Сводка', 'Обзор'), ('Разыскиваются', 'Разыскиваются')],
                      ['Новые IP — впервые замеченные в этот день; в «Итого» — всего разных IP ботов за период.', 'Заявки — отправки форм ботами; «Принято» — сайт принял их как настоящие.'],
                      kpi_col='Визиты ботов|Спам форм', outages=E.get('сбои'))
    # Разыскиваются
    wanted(wb, Pf)
    # Сигнатуры
    Sg = Pf.get('сигнатуры')
    if Sg is not None and len(Sg):
        nm = 'Сигнатуры'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'ID': Sg['id'], 'Правило': Sg['правило'], 'Главное обвинение': Sg['вид'], 'Дело': Sg['дело'], 'Кличка': Sg['кличка'], 'IP': Sg['IP'].astype(int),
                          'Запросов': Sg['запросов'].astype(int), 'Проверка на логе': Sg['проверка'], 'Задевает людей': Sg['ложных'].astype(int)})
        st_ = res.get('stix') or {}
        if st_:   # код правила в STIX — по нему правило находится в файле поиском
            from .stix import sid
            in_ = set(st_.get('сигнатуры_в_пакете') or [])
            d['Код в STIX'] = [sid('indicator', f'sigma:{x}') if x in in_ else '' for x in d['ID']]   # пусто — у дела «к сведению» нет адресов, правило не выгружается
        data_sheet(wb, nm, d, nm, 'Правила, по которым узнаётся каждый атакующий профиль; ID постоянный — один и тот же для этого правила на любом сайте' +
                   (f". Машинный вид — файл {st_.get('файл')} (STIX 2.1) в архиве отчёта" if st_ else ''),
                   {'ID': 22, 'Правило': 50, 'Главное обвинение': 26, 'Кличка': 30, 'Проверка на логе': 50}, wrap=('Правило', 'Главное обвинение', 'Кличка', 'Проверка на логе'),
                   kpi=[('Сигнатур', len(d)), ('Без задетых людей', int((d['Задевает людей'] == 0).sum()))], kpi_col='Запросов',
                   row_rule=lambda r: F_GREY if r.get('Задевает людей') else None, red=(), links=[('Карточки', 'Разыскиваются')])
        notes(wb[nm], ['ID — хеш самого правила: одинаковое правило в любой проверке и на любом сайте получает один и тот же ID. Справочник сигнатур — data/reference/learned/signatures.json.',
                       'Задевает людей — сколько визитов людей пришло из тех же сетей: если не ноль, сеть целиком закрывать нельзя, только по поведению.'])
        legend(wb[nm], ['задевает'])
    # Состав дел
    Cm = Pf.get('состав')
    if Cm is not None and len(Cm):
        nm = 'Состав дел'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Дело': Cm['дело'], 'Кличка': Cm['кличка'], 'IP': Cm['ip'], 'Подсеть': Cm['подсеть'], 'Сеть': Cm['сеть'], 'Страна': Cm['страна'], 'Роль': Cm['роль'],
                          'Запросов': Cm['запросов'].astype(int), 'Первый': Cm['первый'], 'Последний': Cm['последний']})
        data_sheet(wb, nm, d, nm, 'Все IP каждого дела: роль, сеть и активность. Фильтр по делу покажет всю группу, фильтр по подсети — к каким делам она относится',
                   {'Кличка': 34, 'IP': 16, 'Подсеть': 18, 'Сеть': 30, 'Роль': 30, 'Первый': 16, 'Последний': 16}, wrap=('Кличка', 'Сеть', 'Роль'), center=('Страна',),
                   kpi=[('IP', len(d)), ('Дел', int(d['Дело'].nunique())), ('Подсетей', int(d['Подсеть'].nunique()))], kpi_col='Запросов', red=(), links=[('Карточки', 'Разыскиваются')])
    BS = S
    # Операторы
    O, OU = BS.get('Операторы'), BS.get('Операторы: улики')
    if O is not None and len(O):
        nm = 'Операторы'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Оператор': O['оператор'], 'Улики': O['улики'], 'Слабых связей': O['слабые_связи_с_другими'], 'IP спама': O['IP_спама'], 'IP разведки': O['IP_разведки'],
                          'Сети': O['сети'], 'Отправок': O['отправок'], 'Принято': O['принято'], 'Битые входы': O['битые_входы'].astype(str).str.replace(', ', '\n'),
                          'Дней': O['дни'].astype(str).str.count(',') + 1})
        data_sheet(wb, nm, d, nm, 'Связанные группы IP: один человек или программа за разными адресами', {'Оператор': 18, 'Улики': 50, 'Сети': 34, 'Битые входы': 50},
                   wrap=('Улики', 'Сети', 'Битые входы'), kpi=[('Операторов', len(d)), ('Принято заявок', int(d['Принято'].sum()))], kpi_col='Отправок',
                   row_rule=lambda r: F_NOTE if r.get('Принято') else None, red=(), links=[('Карточки', 'Разыскиваются')])
        if OU is not None and len(OU):
            extra_table(wb[nm], pd.DataFrame({'IP A': OU['IP_A'], 'IP B': OU['IP_B'], 'Сила': OU['сила'], 'Улика': OU['улика'], 'Доказательство': OU['доказательство'], 'Учтена': OU['учтена']}),
                        'Улики связей', 'Почему два IP считаются одним оператором', wrap=('Улика', 'Доказательство', 'Учтена'))
    # Спам форм
    SP = BS.get('Спам форм: визиты')
    if SP is not None and len(SP):
        nm = 'Спам форм'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Начало': pd.to_datetime(SP['начало']).dt.strftime('%d.%m.%Y %H:%M'), 'IP': SP['ip'], 'Приём': SP['класс'].astype(str).str.replace('спам форм: ', '', regex=False),
                          'Вход': SP['entry'], 'Код входа': SP['entry_status'].astype(int), 'Страниц': SP['n_pages'].astype(int), 'Файлов': SP['n_static'].astype(int),
                          'Отправок': SP['n_goal'].astype(int), 'Принято': SP['n_conv'].astype(int), 'Сеть': SP['org'].fillna('').astype(str), 'Страна': SP['cc'], 'User-Agent': SP['ua']})
        data_sheet(wb, nm, d, nm, 'Визиты, которые отправили формы как боты', {'Начало': 16, 'IP': 16, 'Приём': 30, 'Вход': 44, 'Сеть': 26, 'User-Agent': 44},
                   wrap=('Приём', 'Вход', 'Сеть', 'User-Agent'), center=('Страна', 'Код входа'), kpi=[('Визитов', len(d)), ('Принято', int(d['Принято'].sum()))], kpi_col='Отправок',
                   row_rule=lambda r: F_NOTE if r.get('Принято') else None, red=())
        legend(wb[nm], ['принята_заявка'])
    # Подделки
    FK = BS.get('Подделки')
    if FK is not None and len(FK):
        nm = 'Подделки'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'IP': FK['ip'], 'Представлялся': FK['представлялся'], 'Запросов': FK['запросов'].astype(int), 'Визитов': FK['визитов'].astype(int),
                          'Сеть': FK['org'].fillna('').astype(str), 'Тип сети': FK['сеть'], 'Страна': FK['страна'], 'Первый': FK['первый'].map(_d), 'Последний': FK['последний'].map(_d),
                          'User-Agent': FK['ip'].astype(str).map(res.get('ua_ip') or {}).fillna('')})   # чем именно бот притворяется
        data_sheet(wb, nm, d, nm, 'Называют себя поисковыми и другими роботами, а приходят не из их сетей', {'IP': 16, 'Представлялся': 24, 'Сеть': 32, 'Тип сети': 22, 'User-Agent': 50},
                   wrap=('Сеть', 'User-Agent'), center=('Страна',), kpi=[('IP', len(d)), ('Запросов', int(d['Запросов'].sum()))], kpi_col='Визитов', red=())
    # Виды ботов
    Sig = A.get('сигнатуры')
    if Sig is not None and len(Sig):
        nm = 'Виды ботов'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Вид': Sig['вид'], 'Признак': Sig['сигнатура'], 'Визитов': Sig['визитов'].astype(int), 'Уникальных IP': Sig['IP'].astype(int), 'Сетей': Sig['сетей'].astype(int),
                          'Дней': Sig['дней'].astype(int), 'Заявок': Sig['заявок'].astype(int), 'Чаще всего из сети': Sig['сеть'], 'Частый вход': Sig['вход']})
        data_sheet(wb, nm, d, nm, 'По какому признаку бот пойман: маскировка, подделка, зонды, спам форм и другие', {'Вид': 18, 'Признак': 40, 'Чаще всего из сети': 26, 'Частый вход': 44},
                   wrap=('Признак', 'Частый вход'), kpi=[('Визитов', int(d['Визитов'].sum()))], kpi_col='Сетей', red=())
        X = BS.get('Явные боты')
        if X is not None and len(X):
            extra_table(wb[nm], pd.DataFrame({'User-Agent': X['ua'], 'Визитов': X['визитов'].astype(int), 'Уникальных IP': X['IP'].astype(int), 'Запросов': X['запросов'].astype(int),
                                              'Сети': X['сети']}).head(200), 'Явные боты', 'Программы, которые не притворяются браузером, — по User-Agent', wrap=('User-Agent', 'Сети'))
    # Боты в рекламе
    AD = EX.get('реклама')
    if AD is not None and len(AD):
        nm = 'Боты в рекламе'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Реклама': AD['channel_sub'], 'Вид бота': AD['вид'], 'Визитов': AD['визитов'].astype(int), 'Уникальных IP': AD['IP'].astype(int), 'Частый вход': AD['вход']})
        data_sheet(wb, nm, d, nm, 'Визиты ботов по рекламным ссылкам — клики, за которые заплачено', {'Реклама': 26, 'Вид бота': 22, 'Частый вход': 60}, wrap=('Частый вход',),
                   kpi=[('Визитов ботов', int(d['Визитов'].sum())), ('Визитов людей с рекламы', int(EX.get('реклама_люди') or 0))], kpi_col='Частый вход', red=(),
                   links=[('Реклама по кампаниям', ('NXLD_05_Marketing.xlsx', 'Кампании'))])
        notes(wb[nm], ['Рекламная система списывает деньги за каждый такой клик. Сети и IP из карточек можно исключить в настройках рекламы и сообщить о скликивании в поддержку.'])
    # Расписание ботов
    H = EX.get('расписание') or {}
    if H:
        heat_sheet(wb, 'Расписание ботов', 'Расписание ботов', 'Когда работают боты: запросов за каждый час по видам',
                   [(k, '', v) for k, v in H.items()], intro=['Куда смотреть: тёмные полосы в одни и те же часы — работа по расписанию; яркие пятна ночью — волны сканирования.'],
                   links=[('Разыскиваются', 'Разыскиваются')])
    # Роботы
    RF, RN = BS.get('Роботы: семейства'), BS.get('Роботы: сети')
    if RF is not None and len(RF):
        nm = 'Роботы'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Робот': RF['семейство'], 'Категория': RF['категория'], 'Запросов': RF['запросов'].astype(int), 'Уникальных IP': RF['IP'].astype(int), 'Дней': RF['дней'].astype(int),
                          'Трафик, МБ': RF['МБ'], 'Ответов 200, %': RF['доля_200_%'], '404': RF['404'].astype(int), '5xx': RF['5xx'].astype(int),
                          'Подлинных, %': RF['подлинных_%'], 'Что смотрел': RF['что_смотрел'].map(nlist)})
        data_sheet(wb, nm, d, nm, 'Роботы, которые честно себя называют: поисковики, сервисы, SEO-роботы', {'Робот': 30, 'Категория': 22, 'Что смотрел': 50}, wrap=('Робот', 'Что смотрел'),
                   kpi=[('Роботов', len(d)), ('Запросов', int(d['Запросов'].sum()))], kpi_col='Дней', red=())
        if RN is not None and len(RN):
            extra_table(wb[nm], pd.DataFrame({'Робот': RN['семейство'], 'Сеть': RN['сеть'].fillna('').astype(str), 'ASN': RN['asn'], 'Запросов': RN['запросов'].astype(int)}).head(200),
                        'Сети роботов', 'Из каких сетей приходит каждый робот', wrap=('Сеть',))
    # Системы мониторинга и утилиты
    M = A.get('мониторинг')
    if M is not None and len(M):
        nm = 'Системы мониторинга'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        sp = M['коды'].map(split_codes)
        d = pd.DataFrame({'Система': M['система'], 'Как опознана': M['как'], 'Что проверяет': M['проверяет'], 'Уникальных IP': M['IP'].astype(int),
                          'Интервал, с': M['интервал'], 'Запросов': M['запросов'].astype(int), 'Ответы': sp.str[0], 'Ошибки': sp.str[1],
                          'User-Agent': M['ua'] if 'ua' in M else ''})   # по UA видно, чей это сервис
        data_sheet(wb, nm, d, nm, 'Сервисы, которые проверяют, работает ли сайт: по имени в User-Agent или по ритму', {'Система': 30, 'Как опознана': 20, 'Что проверяет': 40, 'Ответы': 26, 'Ошибки': 22, 'User-Agent': 50},
                   wrap=('Система', 'Что проверяет', 'Ответы', 'Ошибки', 'User-Agent'), kpi=[('Систем', len(d))], kpi_col='Запросов')
    U = A.get('утилиты')
    if U is not None and len(U):
        nm = 'Утилиты'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Утилита': U['утилита'], 'Запросов': U['запросов'].astype(int), 'Уникальных IP': U['IP'].astype(int), 'Зондов': U['зондов'].astype(int),
                          'Что запрашивали чаще всего': U['что'].astype(str).str.replace(', ', '\n'), 'Похоже на (по IP)': U['похоже'].astype(str).str.replace(', ', '\n')})
        data_sheet(wb, nm, d, nm, 'Программы без браузера: curl, wget, Python, PHP и другие — вывод по поведению каждого IP', {'Утилита': 22, 'Что запрашивали чаще всего': 50, 'Похоже на (по IP)': 40},
                   wrap=('Что запрашивали чаще всего', 'Похоже на (по IP)'), kpi=[('Утилит', len(d))], kpi_col='Зондов', red=())
    ai_sheet(wb, BS.get('ИИ-роботы'), BS.get('ИИ: страницы по запросам людей'))
    # Проверка IP и сети ботов
    CI = BS.get('Проверка IP')
    if CI is not None and len(CI):
        nm = 'Подозреваемые'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        case_of = {}
        for x in Pf.get('дела') or []:
            for ip in x['ips']: case_of[ip] = x['дело']
        d = pd.DataFrame({'IP': CI['ip'], 'Дело': CI['ip'].astype(str).map(case_of).fillna('—'), 'В логе': CI['в_логе'], 'Сеть': CI['сеть'], 'Страна': CI['страна'],
                          'Тип сети': CI['тип_сети'], 'Визитов': CI['визитов'], 'Запросов': CI['запросов'], 'Кто по логу': CI['группа'], 'Отправок форм': CI['отправок'],
                          'Принято': CI['принято'], 'Первый': CI['первый'].map(lambda v: pd.to_datetime(v).strftime('%d.%m.%Y %H:%M') if pd.notna(v) and v else ''),
                          'Последний': CI['последний'].map(lambda v: pd.to_datetime(v).strftime('%d.%m.%Y %H:%M') if pd.notna(v) and v else ''), 'Входы': CI['входы'], 'User-Agent': CI['UA']})
        data_sheet(wb, nm, d, nm, 'IP, которые прислали на проверку: заходили ли они, откуда, что делали и в каком деле', {'IP': 16, 'Сеть': 26, 'Тип сети': 20, 'Кто по логу': 34, 'Входы': 34, 'User-Agent': 40, 'Первый': 16, 'Последний': 16},
                   wrap=('Сеть', 'Кто по логу', 'Входы', 'User-Agent'), center=('В логе', 'Страна', 'Дело'), kpi=[('Прислано', len(d)), ('Нашлись в логе', int((d['В логе'] == 'да').sum())), ('В делах', int((d['Дело'] != '—').sum()))],
                   kpi_col='Запросов', red=(), links=[('Карточки', 'Разыскиваются')])
    MI = Pf.get('меры_ip')
    if MI is not None and len(MI):
        nm = 'Меры по IP'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'IP': MI['ip'], 'Приговор': MI['приговор'], 'Основание': MI['основание'], 'Дело': MI['дело'], 'Кличка': MI['кличка'], 'Подсеть': MI['подсеть'],
                          'Сеть': MI['сеть'], 'Страна': MI['страна'], 'Тип сети': MI['тип_сети'], 'Визитов': MI['визитов'].astype(int), 'Запросов': MI['запросов'].astype(int),
                          'Отправок форм': MI['отправок'].astype(int), 'Первый': MI['первый'], 'Последний': MI['последний']})
        fills = {'заблокировать': F_NOTE, 'ограничить частоту': None, 'наблюдать': None, 'не трогать': F_GREY}
        cnt = d['Приговор'].value_counts()
        data_sheet(wb, nm, d, nm, 'Что делать с каждым IP: готовый список для администратора сервера', {'IP': 16, 'Приговор': 18, 'Основание': 50, 'Кличка': 30, 'Подсеть': 18, 'Сеть': 28, 'Тип сети': 20, 'Первый': 16, 'Последний': 16},
                   wrap=('Основание', 'Кличка', 'Сеть'), center=('Страна', 'Дело'),
                   kpi=[('Заблокировать', int(cnt.get('заблокировать', 0))), ('Ограничить частоту', int(cnt.get('ограничить частоту', 0))), ('Наблюдать', int(cnt.get('наблюдать', 0))), ('Не трогать', int(cnt.get('не трогать', 0)))],
                   kpi_col='Визитов', row_rule=lambda r: fills.get(r.get('Приговор')), red=(), links=[('Карточки дел', 'Разыскиваются')])
        notes(wb[nm], ['Заблокировать — дело серьёзное и из этой сети людей нет. Ограничить частоту — из сети ходят люди или это мобильный оператор (за одним IP много абонентов): закрывать по поведению, не сетью.',
                       'Наблюдать — дело к сведению. Не трогать — сотрудники, сервер сайта и системы мониторинга.'])
        legend(wb[nm], ['приговор_блок', 'свои_норма'])
    BN = BS.get('Бот-сети')
    if BN is not None and len(BN):
        nm = 'Сети ботов'
        if nm not in wb.sheetnames: wb.create_sheet(nm)
        d = pd.DataFrame({'Сеть': BN['сеть'].fillna('').astype(str), 'ASN': BN['asn'], 'Тип сети': BN['nettype'], 'Визитов': BN['визитов'].astype(int), 'Уникальных IP': BN['IP'].astype(int),
                          'Виды': BN['подгруппы'], 'Что делать с адресами': BN['что_делать_с_адресами']})
        data_sheet(wb, nm, d, nm, 'Из каких сетей приходят боты и можно ли закрывать сеть целиком', {'Сеть': 32, 'Тип сети': 22, 'Виды': 50, 'Что делать с адресами': 26}, wrap=('Сеть', 'Виды'),
                   kpi=[('Сетей', len(d))], kpi_col='Визитов', red=())


def overview(wb, res, site):
    Pf = res.get('profiles') or {}
    EX = res.get('bots_extra') or {}
    A = res.get('actors') or {}
    S = (res.get('sheets') or {}).get('Боты', {})
    sm = (res.get('summary') or {}).get('Боты', {})
    if 'Обзор' in wb.sheetnames: del wb['Обзор']
    ws = wb.create_sheet('Обзор', 0)
    W = Wide(ws)
    brand_header(ws, W, res, f'NX LOG DETECTIVE — БОТЫ НА САЙТЕ {site.upper()}', 'Кто атакует сайт, в чём мы его обвиняем и как его узнать', last='I')
    names = set(wb.sheetnames)
    D = Pf.get('дела') or []
    W.r += 1
    W.kpis([('Визитов ботов', int(sm.get('Визитов ботов', 0))), ('IP ботов', int(sm.get('IP ботов', 0))), ('Дел', len(D)),
            ('Срочных', sum(1 for x in D if x['важность'] == 'Срочно')), ('Подделок (IP)', int(sm.get('Подделок роботов (IP)', 0))),
            ('Принято заявок от ботов', int((res.get('сводки') or {}).get('принято_от_ботов', 0)))])   # тот же источник, что журнал и карточки (derive)
    st_ = res.get('stix') or {}
    if st_:   # выгрузка STIX — заметно, сразу под ключевыми цифрами
        W.r += 1
        W.note(f"Сигнатуры, адреса и связи выгружены в STIX 2.1: {st_.get('сигнатур', 0)} сигнатур, {st_.get('дел', 0)} дел, {st_.get('IP', 0)} IP — "
               f"файл {st_.get('файл', '*.stix.json')} в архиве отчёта. STIX — открытый стандарт обмена данными об угрозах; его принимают OpenCTI, MISP и другие платформы.")
        W.r += 1
    if D:
        W.section('Разыскиваются', 'Самые серьёзные дела: главное обвинение, состав и ссылка на карточку.')
        top = D[:10]
        rows = [[f"Дело {x['дело']}", x['кличка'], x['обвинения'][0]['статья'], x['состав']['IP'], x['важность']] for x in top]
        fl = {}
        for k, x in enumerate(top):
            col_ = 'C00000' if x['важность'] == 'Тревога' else (ORANGE if x['важность'] == 'Срочно' else (F_NOTE if x['важность'] == 'Важно' else None))
            if col_: fl[(k, 4)] = col_
        W.table(['Дело', 'Кличка', 'Главное обвинение', 'IP', 'Важность'], rows, ['B', 'CD', 'EFG', 'H', 'I'], num=(3,), center=(4,), wrap=0.95, fills=fl)
        W.link(f'Все дела ({len(D)})', 'Разыскиваются', names)
    Sig = A.get('сигнатуры')
    if Sig is not None and len(Sig):
        W.section('Виды ботов', 'По какому признаку бот пойман.')
        W.table(['Вид', 'Признак', 'Визитов', 'Уникальных IP', 'Заявок'], [[r['вид'], r['сигнатура'], int(r['визитов']), int(r['IP']), int(r['заявок'])] for _, r in Sig.head(8).iterrows()],
                ['B', 'CDE', 'F', 'G', 'H'], num=(2, 3, 4), wrap=0.95)
        W.link('Подробно', 'Виды ботов', names)
    FK = S.get('Подделки')
    if FK is not None and len(FK):
        W.section('Подделки', 'Называют себя известными роботами, а приходят не из их сетей.')
        g = (res.get('сводки') or {}).get('подделки_по_имени')
        if g is None: g = FK.groupby('представлялся').agg(IP=('ip', 'nunique'), запросов=('запросов', 'sum')).sort_values('запросов', ascending=False).reset_index()
        W.table(['Представлялся', 'Уникальных IP', 'Запросов'], [[r['представлялся'], int(r['IP']), int(r['запросов'])] for _, r in g.head(6).iterrows()], ['BC', 'D', 'E'], num=(1, 2))
        W.link('Все подделки', 'Подделки', names)
    O = S.get('Операторы')
    if O is not None and len(O):
        W.section('Спам форм', 'Связанные группы IP, которые отправляют формы.')
        W.table(['Оператор', 'IP спама', 'IP разведки', 'Отправок', 'Принято'], [[r['оператор'], int(r['IP_спама']), int(r['IP_разведки']), int(r['отправок']), int(r['принято'])] for _, r in O.iterrows()],
                ['BC', 'D', 'E', 'F', 'G'], num=(1, 2, 3, 4))
        W.link('Операторы и улики', 'Операторы', names)
    AD = EX.get('реклама')
    if AD is not None and len(AD):
        W.section('Боты в рекламе', 'Визиты ботов по рекламным ссылкам — клики, за которые заплачено.')
        g = (res.get('сводки') or {}).get('реклама_по_системам')
        if g is None: g = AD.groupby('channel_sub').agg(визитов=('визитов', 'sum'), IP=('IP', 'sum')).sort_values('визитов', ascending=False).reset_index()
        W.table(['Реклама', 'Визитов ботов', 'Уникальных IP'], [[r['channel_sub'], int(r['визитов']), int(r['IP'])] for _, r in g.head(6).iterrows()], ['BC', 'D', 'E'], num=(1, 2))
        W.link('Подробно', 'Боты в рекламе', names)
    AI = S.get('ИИ-роботы')
    if AI is not None and len(AI):
        W.section('ИИ-роботы', 'Роботы нейросетей: что читают и как часто.')
        W.table(['Робот', 'Назначение', 'Запросов', 'Подлинных, %'], [[r['робот'], r['назначение'], int(r['запросов']), r['подлинных_%'] if pd.notna(r['подлинных_%']) else '—'] for _, r in AI.head(6).iterrows()],
                ['BC', 'DEF', 'G', 'H'], num=(2,), wrap=0.95)
        W.link('Все ИИ-роботы', 'ИИ-роботы', names)
    W.related(res, 'Боты')
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build(wb, res, S, site):
    for k_ in ('Сводка', 'О данных'):
        if k_ in wb.sheetnames: del wb[k_]
    build_sheets(wb, res, S)
    from . import report_load
    report_load.scanner_ips_sheet(wb, res, name='IP сканеров')
    overview(wb, res, site)
    byname = {w.title: w for w in wb._sheets}
    head_ = [byname[n_] for n_ in ORDER if n_ in byname]
    wb._sheets = head_ + [w for w in wb._sheets if w not in head_]
