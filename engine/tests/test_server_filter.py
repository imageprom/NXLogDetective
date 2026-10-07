"""Серверный фильтр: кому защита сервера отказывает — поисковым роботам, роботам проверки объявлений, людям (лист 03 и карточки)."""
import os
import pickle
import random
import subprocess
import sys
from datetime import datetime, timedelta

import pytest

ENGINE = os.path.join(os.path.dirname(__file__), '..')
sys.path.insert(0, ENGINE)
from nxld.security import waves  # noqa: E402

UA_CHROME = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36'
UA_GOOGLE = 'Mozilla/5.0 (compatible; Googlebot/2.1; +http://www.google.com/bot.html)'
UA_BING = 'Mozilla/5.0 (compatible; bingbot/2.0; +http://www.bing.com/bingbot.htm)'
UA_DIRECT = 'Mozilla/5.0 (compatible; YaDirectFetcher/1.0; Dyatel; +http://yandex.com/bots)'
PAGES = ['/', '/catalog/', '/catalog/?PAGEN_1=2', '/catalog/shary/?sort=price', '/catalog/shary/', '/catalog/girlyandy/', '/about/', '/contacts/', '/dostavka/', '/akcii/']
ASSETS = ['/local/templates/main/style.css', '/local/templates/main/script.js', '/upload/iblock/1.jpg']
WAVE = range(4, 7)   # дни 4–6 сентября: фильтр включён


def _robot_ip(family, n):
    """IP из официальной сети робота (data/reference/robot_ranges.json)."""
    import ipaddress, json
    d = json.load(open(os.path.join(ENGINE, '..', 'data', 'reference', 'robot_ranges.json'), encoding='utf-8'))['семейства'][family]
    net = ipaddress.ip_network(next(x for x in d if x.startswith('66.249.')) if family == 'Googlebot' else d[0])
    return [str(h) for _, h in zip(range(n), net.hosts())]


def _log(path):
    random.seed(7)
    t0 = datetime(2026, 9, 1)
    L = []
    def line(ip, t, url, code, ua, ref='-', size=None):
        L.append(f'{ip} - - [{t:%d/%b/%Y:%H:%M:%S} +0300] "GET {url} HTTP/1.1" {code} {size or random.randint(2000, 60000)} "{ref}" "{ua}"')
    def visit(ip, t, code_after_first=200):
        p = random.choice(PAGES)
        line(ip, t, p, 200, UA_CHROME, 'https://yandex.ru/')
        for a in ASSETS: line(ip, t + timedelta(seconds=1), a, 200, UA_CHROME, f'https://site.ru{p}')
        for k in range(3):
            line(ip, t + timedelta(seconds=30 * (k + 1)), random.choice(PAGES), code_after_first, UA_CHROME, f'https://site.ru{p}')
    gip, bip = _robot_ip('Googlebot', 3), _robot_ip('Bingbot', 2)
    ru = ['95.165.10.20', '95.165.11.30', '176.59.40.10', '94.25.170.10', '46.138.1.1']
    foreign = ['81.2.69.142', '104.244.72.10', '212.58.244.20']
    for day in range(1, 11):
        base = t0 + timedelta(days=day - 1)
        blocked = day in WAVE
        for k in range(60):   # Googlebot
            line(random.choice(gip), base + timedelta(minutes=20 * k + 1), random.choice(PAGES), 403 if blocked else 200, UA_GOOGLE)
        for k in range(20):   # Bingbot
            line(random.choice(bip), base + timedelta(minutes=60 * k + 7), random.choice(PAGES), 403 if blocked else 200, UA_BING)
        for k in range(6):    # робот проверки объявлений Яндекс Директа
            line('5.255.253.10', base + timedelta(minutes=200 * k + 3), '/catalog/shary/', 403 if blocked else 200, UA_DIRECT)
        for k in range(14):   # люди из России — всегда 200
            visit(ru[k % len(ru)], base + timedelta(hours=8, minutes=50 * k))
        for k in range(5):    # люди из-за рубежа — в волну отказ после первой страницы
            visit(foreign[k % len(foreign)], base + timedelta(hours=9, minutes=70 * k + 5), 403 if blocked else 200)
    L.sort(key=lambda s: datetime.strptime(s.split('[')[1][:20], '%d/%b/%Y:%H:%M:%S'))
    open(path, 'w').write('\n'.join(L) + '\n')
    return len(L)


@pytest.fixture(scope='module')
def run(tmp_path_factory):
    d = tmp_path_factory.mktemp('sf')
    log, work, out = d / 'access.log', d / 'work', d / 'out'
    n = _log(log)
    assert 2000 <= n <= 3000, n
    learned = os.path.join(ENGINE, '..', 'data', 'reference', 'learned')   # прогон дописывает справочник — после теста вернуть как было
    keep = {os.path.join(r, f): open(os.path.join(r, f), 'rb').read() for r, _, fs in os.walk(learned) for f in fs}
    try:
        p = subprocess.run([sys.executable, os.path.join(ENGINE, 'nxld_run.py'), '--logs', str(log), '--work', str(work), '--out', str(out)],
                           capture_output=True, text=True, timeout=600)
    finally:
        for r, _, fs in os.walk(learned):
            for f in fs:
                if os.path.join(r, f) not in keep: os.remove(os.path.join(r, f))
        for f, b in keep.items(): open(f, 'wb').write(b)
    assert p.returncode == 0, p.stdout[-2000:] + p.stderr[-4000:]
    assert 'Traceback' not in p.stdout + p.stderr, p.stdout[-2000:] + p.stderr[-4000:]
    res = pickle.load(open(work / 'results.pkl', 'rb'))
    return res, out


def _cells(path, sheet):
    import openpyxl
    ws = openpyxl.load_workbook(path)[sheet]
    return [v for row in ws.iter_rows(values_only=True) for v in row if v is not None]


def test_waves():
    assert waves(['2026-09-01', '2026-09-02', '2026-09-18', '2026-09-19', '2026-09-24', '2026-09-20', '2026-09-21', '2026-09-22', '2026-09-23']) == '1–2.09, 18–24.09'
    assert waves(['2026-09-30', '2026-10-01', '2026-10-05']) == '30.09–1.10, 5.10'
    assert waves([]) == ''


def test_sheet_in_03_with_three_groups(run):
    import openpyxl
    res, out = run
    wb = openpyxl.load_workbook(out / 'NXLD_03_Load_Security.xlsx')
    assert 'Серверный фильтр' in wb.sheetnames
    names = wb.sheetnames
    assert names.index('Серверный фильтр') < names.index('Сканеры') if 'Сканеры' in names else True
    vals = _cells(out / 'NXLD_03_Load_Security.xlsx', 'Серверный фильтр')
    for g in ('Поисковые роботы', 'Роботы проверки объявлений', 'Люди'):
        assert g in vals, g
    for who in ('Googlebot', 'Bingbot', 'YaDirectFetcher'):
        assert who in vals, who
    assert '4–6.09' in vals
    D = res['security']['фильтр']['таблица']
    assert set(D['группа']) == {'Поисковые роботы', 'Роботы проверки объявлений', 'Люди'}
    assert not (D.loc[D['группа'] == 'Люди', 'страна'] == 'RU').any()   # российским людям не отказывали
    ov = _cells(out / 'NXLD_03_Load_Security.xlsx', 'Обзор')
    assert 'СЕРВЕРНЫЙ ФИЛЬТР' in ov   # блок Обзора (заголовки блоков — прописными)


def test_raw_sheet_gone_from_02(run):
    import openpyxl
    _, out = run
    wb = openpyxl.load_workbook(out / 'NXLD_02_Errors.xlsx')
    assert not any(n.startswith('Защита') for n in wb.sheetnames)
    assert 'Защита: кого блокирует' not in run[0]['sheets'].get('Ошибки', {})


def test_cards(run):
    res, _ = run
    by = {x['key'].split(':')[1]: x for x in res['findings']}
    s = by['search_blocked']
    assert (s['блок'], s['важность'], s['лист']) == ('SEO', 'Срочно', 'Серверный фильтр')
    assert 'Googlebot — ' in s['факты'] and 'Bingbot — ' in s['факты'] and '4–6.09' in s['факты']
    a = by['ad_checker_blocked']
    assert (a['блок'], a['важность']) == ('Маркетинг', 'Срочно')
    assert 'YaDirectFetcher' in a['факты']
    p = by['blocked_people']
    assert p['блок'] == 'Нагрузка и безопасность'
    assert p['key'] == 'Нагрузка и безопасность:blocked_people:all'
    assert 'Исключить из фильтра подлинных поисковых роботов' in s['что_сделать']
