"""#17: «Дней» у оператора — число разных дней активности, а не число запятых в списке, обрезанном до 20."""
import os
import pickle
import tempfile
from datetime import timedelta

import openpyxl

import synth
from nxld import analyze, prepare, report


def _log(days=25):
    log = synth.Log()
    log.people(days=days)
    for d in range(days):
        t = synth.T0 + timedelta(days=d, hours=9)
        log.visit(synth.RU[0], t)   # человек, который отправляет форму
        log.line(synth.RU[0], t + timedelta(seconds=100), '/contacts/send/', 200, ref='https://site.ru/contacts/', method='POST')
        for k, ip in enumerate(('185.100.0.1', '185.100.0.2')):   # спам: прямой заход на битый адрес → главная → форма, два IP на один вход
            t2 = synth.T0 + timedelta(days=d, hours=14, minutes=10 * k)
            log.line(ip, t2, '/zz-dead/', 404, ref='-')
            log.line(ip, t2 + timedelta(seconds=2), '/', 200, ref='-')
            log.line(ip, t2 + timedelta(seconds=5), '/contacts/send/', 200, ref='https://site.ru/', method='POST')
    return log


def test_operator_days_not_cut_at_20(monkeypatch):
    monkeypatch.setenv('NXLD_STRICT_CATALOG', '1')
    with tempfile.TemporaryDirectory() as d:
        path, work, out = os.path.join(d, 'access.log'), os.path.join(d, 'work'), os.path.join(d, 'out')
        _log().write(path)
        prepare.run([path], work, log=lambda *a: None)
        res = analyze.run(work, log=lambda *a: None)
        O = res['sheets']['Боты']['Операторы']
        assert len(O) and int(O['дней'].max()) == 25, O.get('дней')   # без обрезки
        assert O['дни'].iloc[0].count(',') + 1 == 20   # текст списка по-прежнему обрезан
        report.build(res, out, None, None, os.path.join(work, 'NXLD_Redmine.textile'), None, work, False)
        ws = openpyxl.load_workbook(os.path.join(out, 'NXLD_04_Bots.xlsx'))['Операторы']
        hdr = {c.value: c.column for c in ws[4] if c.value}
        assert max(ws.cell(r, hdr['Дней']).value for r in range(5, 5 + len(O))) == 25
