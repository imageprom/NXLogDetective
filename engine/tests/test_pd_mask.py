"""Маскировка персональных данных: рекламные и поисковые идентификаторы — не ПД, телефон — только похожее на телефон."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from nxld.recon import has_pd, mask_pd, is_phone, is_ad_key  # noqa: E402


def test_clid_not_pd():
    q = 'clid=2270455-308&text=шары'
    assert not has_pd(q)
    assert mask_pd(q) == q


def test_yclid_utm_not_pd():
    q = 'yclid=17983564959511740415&utm_source=yandex'
    assert not has_pd(q)
    assert mask_pd(q) == q


def test_phone_masked_clid_kept():
    q = 'phone=%2B7%20(903)%20123-45-67&clid=2270455-308'
    assert has_pd(q)
    m = mask_pd(q)
    assert m.endswith('&clid=2270455-308')
    assert m.startswith('phone=')
    assert '123' not in m and '45' not in m.split('&')[0][:-2]
    assert m.split('&')[0].endswith('67')


def test_email_masked_gclid_kept():
    q = 'email=ivan%40mail.ru&gclid=Cj0KCQ'
    assert has_pd(q)
    m = mask_pd(q)
    assert m.endswith('&gclid=Cj0KCQ')
    assert 'ivan' not in m and 'mail.ru' not in m and 'ivan%40' not in m


def test_long_id_not_pd():
    assert not has_pd('id=12345678901234567890')


def test_tel_pd():
    q = 'tel=89031234567'
    assert has_pd(q)
    assert '1234' not in mask_pd(q)


def test_ad_keys_from_reference():
    for k in ('clid', 'yclid', 'gclid', 'ysclid', 'fbclid', 'rb_clickid', 'vkclid', 'msclkid', 'utm_source', 'utm_term',
              'calltouch_tm', 'etext', '_openstat', 'roistat'):
        assert is_ad_key(k), k
    for k in ('phone', 'tel', 'email', 'name', 'id', 'amp;phone'):
        assert not is_ad_key(k), k


def test_is_phone():
    assert is_phone('+7 (903) 123-45-67')
    assert is_phone('89031234567')
    assert is_phone('+44 20 7946 0958')
    assert not is_phone('2270455-308')
    assert not is_phone('17983564959511740415')
    assert not is_phone('+1234567890123456')


def test_order_and_separators_kept():
    q = '/page/?utm_source=ya&phone=89031234567&page=2&flag&clid=1-2'
    m = mask_pd(q)
    assert m.startswith('/page/?utm_source=ya&phone=')
    assert m.endswith('&page=2&flag&clid=1-2')
    assert '1234' not in m


def test_plain_text_as_before():
    assert mask_pd('звоните 8 903 123-45-67') == 'звоните 8** ***-**-67'
    assert mask_pd('ivan@mail.ru') == 'i***@m***'


def test_checkword_secret_and_email():
    m = mask_pd('login=yes&USER_CHECKWORD=diea4ns7iuco1vzwt7det6ocjiy383mj&USER_LOGIN=a%40b.ru')
    assert m.startswith('login=yes&USER_CHECKWORD=diea…&USER_LOGIN=')
    assert 'diea4ns7' not in m and 'a%40b.ru' not in m and 'a@b.ru' not in m


def test_sessid_secret_not_pd():
    q = 'sessid=1f4e04b22e04cef40044c820f16100e4&clid=2270455-308'
    assert not has_pd(q)
    assert mask_pd(q) == 'sessid=1f4e…&clid=2270455-308'


def test_secret_keys_case_insensitive():
    from nxld.recon import mask_value
    for k in ('token', 'ACCESS_TOKEN', 'api_key', 'ApiKey', 'password', 'passwd', 'checkword', 'PHPSESSID'):
        assert mask_value(k, 'abcdefgh') == 'abcd…', k


def test_mask_cells_secrets_without_pd():
    from openpyxl import Workbook
    from nxld.report_tables import mask_pd_cells
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '/auth/?login=yes&USER_CHECKWORD=diea4ns7iuco1vzwt7det6ocjiy383mj'
    ws['A2'] = 'name=Квартал Заречный'
    mask_pd_cells(wb)
    assert ws['A1'].value == '/auth/?login=yes&USER_CHECKWORD=diea…'
    assert ws['A2'].value == 'name=Квартал Заречный'


def test_params_sheet_top_value():
    from openpyxl import Workbook
    from nxld.report_tables import params_sheet
    def row(k, v):
        return {'параметр': k, 'ключ': k, 'группа': 'Логика сайта', 'запросов': 5, 'людей': 2, 'значений': 1, 'частое_значение': v, 'где': ''}
    wb = Workbook()
    ws = params_sheet(wb, {'params': [row('phone', '89031234567'), row('clid', '2270455-308'), row('USER_CHECKWORD', 'diea4ns7iuco')]})
    vals = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
    assert not any('1234567' in v for v in vals)
    assert any(v.startswith('8') and v.endswith('-67') for v in vals)   # телефон есть, но маскированный
    assert '2270455-308' in vals
    assert 'diea…' in vals and not any('diea4ns7' in v for v in vals)
